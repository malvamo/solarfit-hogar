from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import build_final_report as sf


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = PROJECT_ROOT.parents[2]
OUTPUT = WORKSPACE / "Entregables" / "CalculoDiferencial" / "Guia_de_Uso_y_Pruebas_SolarFit_Hogar.docx"
EVIDENCE_DIR = PROJECT_ROOT / "cases" / "evidence" / "latest"
ASSET_DIR = PROJECT_ROOT / "docs" / "assets"
APP_URL = "https://solarfit-hogar.vercel.app"


def fmt(value: float, decimals: int = 2) -> str:
    text = f"{value:,.{decimals}f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")


def load_cases() -> list[dict]:
    cases = []
    for path in sorted((EVIDENCE_DIR / "cases").glob("*.json")):
        with path.open("r", encoding="utf-8") as stream:
            cases.append(json.load(stream))
    if len(cases) != 5:
        raise RuntimeError(f"Se esperaban 5 casos de evidencia y se encontraron {len(cases)}.")
    return cases


def find_case(cases: list[dict], prefix: str) -> dict:
    for case in cases:
        if case["caseId"].startswith(prefix):
            return case
    raise RuntimeError(f"No se encontró el caso {prefix}.")


def configure_friendly_styles(doc: Document) -> None:
    sf.configure_styles(doc)
    styles = doc.styles

    normal = styles["Normal"]
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(7)

    for name, size in (("Title", 28), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title_properties = title._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)

    caption = styles["Caption"]
    caption.font.color.rgb = RGBColor.from_string(sf.MUTED)


def configure_section(section, *, first_page: bool = False) -> None:
    sf.set_page_geometry(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    text = "Guía interna del equipo · SolarFit Hogar"
    if not first_page:
        text += " · página "
    sf.set_run_font(paragraph.add_run(text), size=8.5, color=sf.MUTED)
    if not first_page:
        sf.add_page_field(paragraph)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(12)

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(18)
    sf.set_run_font(label.add_run("GUÍA INTERNA DEL EQUIPO"), size=10, color=sf.TEAL, bold=True)

    title_paragraph = doc.add_paragraph("Guía sencilla para usar SolarFit Hogar", style="Title")
    title_properties = title_paragraph._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    doc.add_paragraph(
        "Qué construimos, cómo funciona y cómo crear nuestros propios ejemplos",
        style="Subtitle",
    )

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.left_indent = Inches(0.65)
    intro.paragraph_format.right_indent = Inches(0.65)
    intro.paragraph_format.space_after = Pt(26)
    sf.set_run_font(
        intro.add_run(
            "Esta guía está pensada para conversar y practicar entre compañeros. "
            "No es un entregable oficial del proyecto."
        ),
        size=12,
        color=sf.INK,
        bold=True,
    )

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    sf.set_run_font(paragraph.add_run("Aplicación"), size=10, color=sf.MUTED)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sf.set_run_font(paragraph.add_run(APP_URL), size=12, color=sf.BLUE, bold=True)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(34)
    sf.set_run_font(
        paragraph.add_run("Versión para el grupo · septiembre de 2026"),
        size=9.5,
        color=sf.MUTED,
        italic=True,
    )


def add_heading(doc: Document, text: str, *, level: int = 1, page_break: bool = False) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    if page_break:
        paragraph.paragraph_format.page_break_before = True


def add_note(doc: Document, lead: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(9)
    sf.set_run_font(paragraph.add_run(f"{lead} "), bold=True, color=sf.INK)
    sf.set_run_font(paragraph.add_run(text), color=sf.INK)


def add_equation(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    sf.set_run_font(paragraph.add_run(text), size=13, color=sf.INK, bold=True)


def add_image(doc: Document, path: Path, caption: str, alt_text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(6.05))
    shape._inline.docPr.set("descr", alt_text)
    caption_paragraph = doc.add_paragraph(caption, style="Caption")
    caption_paragraph.paragraph_format.keep_with_next = False


def add_quick_start(doc: Document, steps_id: int) -> None:
    add_heading(doc, "Empecemos por lo más importante")
    sf.add_para(
        doc,
        "SolarFit es una calculadora web que nos ayuda a responder una pregunta sencilla: "
        "¿cuántos paneles solares necesitaría un hogar para producir una cantidad de energía parecida a la que consume?",
    )
    sf.add_para(
        doc,
        "La respuesta es una primera estimación. Sirve para comparar escenarios, entender qué cambia el resultado "
        "y detectar si el espacio del techo parece suficiente. No reemplaza la visita de un instalador ni un diseño eléctrico.",
    )

    add_heading(doc, "Prueba rápida de cinco minutos", level=2)
    for text in (
        f"Abra {APP_URL}.",
        "Pulse Cargar caso calibrado.",
        "Entre a la pestaña Datos y confirme que aparecen 156 kWh en cada mes.",
        "Entre a Modelo. Allí verá la fuente solar, el panel y la explicación matemática.",
        "Entre a Resultados. Debe ver 4 paneles y 1,80 kWp.",
    ):
        sf.add_step(doc, text, steps_id)

    add_note(
        doc,
        "Si llegó a ese resultado,",
        "ya hizo la prueba básica. En las páginas siguientes explicamos qué significa y cómo crear un ejemplo diferente.",
    )

    sf.add_table(
        doc,
        ["La aplicación recibe", "La aplicación entrega"],
        [
            ("Consumo de energía de 12 meses", "Cantidad continua y número entero de paneles"),
            ("Ubicación general y recurso solar", "Producción estimada y porcentaje de cobertura"),
            ("Potencia y tamaño del panel", "Espacio aproximado necesario"),
            ("Área útil del techo", "Aviso de techo viable o restringido"),
        ],
        [4480, 4880],
        font_size=9.6,
    )


def add_what_we_built(doc: Document, bullet_id: int) -> None:
    add_heading(doc, "Qué construimos", page_break=True)
    sf.add_para(
        doc,
        "Construimos la primera versión usable de SolarFit Hogar. En tecnología suele llamarse MVP: "
        "una versión que ya funciona y permite probar la idea, aunque todavía no es un producto comercial completo.",
    )
    sf.add_para(doc, "El sistema hace cuatro cosas principales:")
    for text in (
        "organiza doce meses de consumo de energía;",
        "consulta una estimación de producción solar para la ubicación;",
        "calcula cuántos paneles se necesitarían;",
        "compara la recomendación con el espacio disponible del techo.",
    ):
        sf.add_bullet(doc, text, bullet_id)

    add_heading(doc, "Cómo funciona por dentro", level=2)
    sf.add_para(
        doc,
        "Piense en el cálculo como una división. Primero sumamos la energía que queremos cubrir. "
        "Después calculamos cuánta energía produciría un panel. Finalmente dividimos la necesidad entre la producción de un panel.",
    )
    add_equation(doc, "paneles necesarios = energía que queremos cubrir ÷ energía que produce un panel")
    sf.add_para(
        doc,
        "El resultado suele tener decimales. Como no podemos instalar 3,30 paneles, el sistema redondea hacia arriba y recomienda 4. "
        "Luego revisa si esos cuatro paneles caben en el área de techo que escribimos.",
    )

    add_heading(doc, "Dónde entra Cálculo Diferencial", level=2)
    sf.add_para(
        doc,
        "Usamos la derivada para responder algo útil: si cambia el rendimiento solar, ¿qué tan rápido cambia la necesidad de paneles? "
        "El signo negativo indica que, si cada panel produce más energía, se necesitan menos paneles.",
    )
    add_equation(doc, "n(Y) = qD ÷ (P·Y)      n′(Y) = −qD ÷ (P·Y²)")
    add_note(
        doc,
        "Eso es todo lo esencial,",
        "no añadimos otros temas de cálculo que no cambien o expliquen una decisión del sistema.",
    )
    add_image(
        doc,
        ASSET_DIR / "solarfit-modelo.png",
        "La pestaña Modelo muestra la fuente solar, las fórmulas y el efecto del rendimiento.",
        "Pantalla de SolarFit en la pestaña Modelo con fuente solar, fórmulas y gráfica de sensibilidad.",
    )


def add_guided_example(doc: Document, bogota: dict, steps_id: int) -> None:
    result = bogota["results"]["pvgis"]
    pvgis_yield = result["sensitivity"]["referenceYield"]
    annual_bias = bogota["validation"]["metrics"]["specificYield"]["annualBiasPct"]

    add_heading(doc, "Nuestro primer ejemplo con Bogotá", page_break=True)
    sf.add_para(
        doc,
        "El ejemplo inicial no representa una casa específica. Usa un promedio oficial de consumo para piso térmico frío, "
        "un panel comercial de 450 W, un punto de referencia en Bogotá y una serie solar consultada en PVGIS. "
        "Por eso aparece marcado como simulación calibrada y no como medición.",
    )

    for text in (
        "Pulse Cargar caso calibrado para volver siempre al mismo punto de partida.",
        "En Datos revise el consumo, el área del techo y las características del panel.",
        "En Modelo pulse Consultar PVGIS 5.3 si quiere actualizar la estimación solar.",
        "Pulse Contrastar con PVWatts para pedir una segunda estimación independiente.",
        "Abra Resultados y compare las cifras con el resumen siguiente.",
    ):
        sf.add_step(doc, text, steps_id)

    sf.add_table(
        doc,
        ["Dato", "Resultado esperado", "Cómo leerlo"],
        [
            ("Consumo anual", f"{fmt(result['totalConsumptionKwh'], 0)} kWh", "Suma de los doce meses"),
            ("Cantidad continua", fmt(result["continuousPanelCount"], 2), "Resultado antes de redondear"),
            ("Paneles recomendados", str(result["recommendedPanelCount"]), "Cantidad física mínima según el modelo"),
            ("Potencia instalada", f"{fmt(result['recommendedCapacityKwp'], 2)} kWp", "Potencia conjunta de los cuatro paneles"),
            ("Producción anual", f"{fmt(result['buildableAnnualGenerationKwh'], 0)} kWh", "Energía estimada por PVGIS"),
            ("Cobertura anual", f"{fmt(result['buildableAnnualCoveragePct'], 0)} %", "Producción estimada dividida por consumo"),
            ("Máximo por techo", str(result["maximumPanelsByRoof"]), "Límite preliminar basado solo en área"),
        ],
        [2100, 2050, 5210],
        font_size=9.2,
    )
    add_note(
        doc,
        "Comprobación solar,",
        f"PVGIS usa aproximadamente {fmt(pvgis_yield, 0)} kWh/kWp al año. "
        f"PVWatts mantiene la decisión de cuatro paneles y difiere cerca de {fmt(abs(annual_bias), 1)} % en el total anual.",
    )
    add_image(
        doc,
        ASSET_DIR / "solarfit-resultados.png",
        "La pestaña Resultados resume la recomendación y después muestra el detalle de cada mes.",
        "Pantalla de resultados de SolarFit con cuatro paneles, potencia, cobertura anual y gráfica mensual.",
    )


def add_own_case(doc: Document, steps_id: int) -> None:
    add_heading(doc, "Cómo crear nuestro propio caso", page_break=True)
    sf.add_para(
        doc,
        "Podemos construir tantos ejemplos como queramos. Lo importante es dejar claro de dónde salió cada dato. "
        "Un ejercicio inventado sirve para aprender, pero no debe presentarse como una vivienda real.",
    )

    sf.add_table(
        doc,
        ["Tipo de caso", "Datos que usamos", "Cómo debemos nombrarlo"],
        [
            ("Real", "Facturas anonimizadas y área medida", "Caso real anonimizado"),
            ("Simulado calibrado", "Promedio oficial, panel documentado y supuestos visibles", "Simulación calibrada"),
            ("Ejercicio libre", "Valores elegidos solo para probar", "Ejercicio de práctica"),
        ],
        [1800, 4300, 3260],
        font_size=9.3,
    )

    add_heading(doc, "Paso a paso", level=2)
    for text in (
        "Pulse Crear caso propio.",
        "Escriba un nombre que diga si el caso es real, simulado o de práctica.",
        "Use solo una ciudad o ubicación general. No escriba una dirección exacta.",
        "Complete los doce consumos en kWh. Puede escribirlos uno por uno o importar la plantilla CSV.",
        "Indique el área útil del techo. Si es una suposición, anótelo cuando documente el caso.",
        "Copie potencia, ancho y alto desde la ficha del panel que quiera evaluar.",
        "Elija Estimación con PVGIS y pulse Consultar PVGIS 5.3.",
        "Elija Balance anual para comparar el total del año o Meta en cada mes para una condición más exigente.",
        "Abra Resultados, lea las advertencias y exporte el proyecto en JSON o los resultados en CSV.",
    ):
        sf.add_step(doc, text, steps_id)

    add_heading(doc, "Ficha para preparar un ejemplo", level=2)
    sf.add_table(
        doc,
        ["Campo", "Qué escribir", "Mi dato y su origen"],
        [
            ("Nombre", "Ejemplo: práctica de techo pequeño", "____________________________"),
            ("Ubicación", "Ciudad o coordenada general", "____________________________"),
            ("Consumo", "12 valores en kWh", "____________________________"),
            ("Techo útil", "Área en m²", "____________________________"),
            ("Panel", "Modelo, Wp, ancho y alto", "____________________________"),
            ("Meta", "Porcentaje que queremos cubrir", "____________________________"),
            ("Fuente solar", "PVGIS o HSP documentadas", "____________________________"),
        ],
        [1700, 3650, 4010],
        font_size=9.2,
    )


def add_practice_ideas(doc: Document, bullet_id: int) -> None:
    add_heading(doc, "Ideas sencillas para crear ejemplos", page_break=True)
    sf.add_para(
        doc,
        "La forma más fácil de aprender es copiar el caso de Bogotá y cambiar una sola cosa cada vez. "
        "Así sabemos qué variable causó el cambio en el resultado.",
    )

    sf.add_table(
        doc,
        ["Ejemplo", "Qué cambiar", "Qué deberíamos observar"],
        [
            ("Más consumo", "Cambiar cada mes de 156 a 200 kWh", "La necesidad de paneles no debe bajar"),
            ("Techo pequeño", "Reducir el área útil de 20 a 6 m²", "El sistema debe avisar que no caben todos"),
            ("Meta exigente", "Cambiar de balance anual a meta mensual", "La cantidad puede subir para cubrir el mes más difícil"),
            ("Otra ciudad", "Cambiar ubicación y volver a consultar PVGIS", "El rendimiento y la recomendación pueden cambiar"),
            ("Otro panel", "Usar una ficha oficial diferente", "La potencia y el tamaño afectan paneles y área"),
        ],
        [1750, 3400, 4210],
        font_size=9.2,
    )

    add_note(
        doc,
        "Regla para no confundirnos,",
        "si elegimos 200 kWh solo para experimentar, escribimos ejercicio de práctica. Si queremos defenderlo como simulación calibrada, necesitamos una fuente o una regla documentada.",
    )

    add_heading(doc, "Qué aprendemos de los cinco casos incluidos", level=2)
    for text in (
        "Barranquilla muestra un consumo de referencia mayor y buena disponibilidad solar.",
        "Bogotá es nuestro caso base para aprender a usar la aplicación.",
        "Cali y Medellín permiten comparar cómo cambia el recurso solar entre ciudades.",
        "Quibdó muestra un techo que solo admite tres paneles aunque el modelo necesita cuatro.",
        "El contraste PVGIS y PVWatts nos recuerda que una simulación no es una medición exacta.",
    ):
        sf.add_bullet(doc, text, bullet_id)


def add_results_explanation(doc: Document) -> None:
    add_heading(doc, "Cómo entender los resultados", page_break=True)
    sf.add_table(
        doc,
        ["Lo que aparece", "Explicación sencilla"],
        [
            ("Cantidad continua", "El cálculo antes de convertirlo en paneles completos. Por ejemplo, 3,30."),
            ("Paneles recomendados", "El entero redondeado hacia arriba. En el ejemplo, 4."),
            ("Potencia instalada", "La suma de la potencia nominal de todos los paneles."),
            ("Producción anual", "La energía que el modelo estima que produciría el sistema en un año."),
            ("Cobertura", "Producción estimada comparada con consumo. Puede superar 100 %."),
            ("Mes crítico", "El mes en el que la cobertura es menor."),
            ("Máximo por techo", "Cuántos paneles cabrían usando solo el área declarada."),
            ("Rango de incertidumbre", "Cómo puede variar el redondeo si el rendimiento resulta mejor o peor."),
        ],
        [2450, 6910],
        font_size=9.4,
    )

    add_heading(doc, "Dos decisiones que no debemos confundir", level=2)
    add_note(
        doc,
        "Balance anual,",
        "busca que la producción de todo el año compense la meta anual. Es normal que algún mes quede un poco por debajo de 100 %.",
    )
    add_note(
        doc,
        "Meta en cada mes,",
        "elige suficientes paneles para intentar cumplir incluso en el mes más difícil. Por eso suele pedir la misma cantidad o una mayor.",
    )

    add_heading(doc, "PVGIS y PVWatts", level=2)
    sf.add_para(
        doc,
        "PVGIS es la estimación principal. PVWatts funciona como una segunda opinión. Si dan valores distintos, no significa que uno esté dañado: "
        "usan bases meteorológicas y métodos diferentes. Debemos mostrar la diferencia y evitar prometer que el resultado será exacto.",
    )


def add_test_plan(doc: Document, steps_id: int) -> None:
    add_heading(doc, "Cómo probar que todo funciona", page_break=True)
    sf.add_para(
        doc,
        "Cada compañero puede hacer estas pruebas. No es necesario saber programar. Basta con cambiar datos, mirar la respuesta y marcar el resultado observado.",
    )
    tests = (
        "Cargue Bogotá y confirme 4 paneles, 1,80 kWp y cerca de 121 % de cobertura anual.",
        "Suba el consumo y confirme que la necesidad de paneles no disminuye.",
        "Reduzca el techo y confirme que el sistema separa paneles necesarios de paneles que caben.",
        "Cambie a Meta en cada mes y compare la cantidad con Balance anual.",
        "Consulte PVGIS y después PVWatts. Confirme que aparecen las dos estimaciones.",
        "Exporte el proyecto JSON. Ábralo o vuelva a importarlo para comprobar que conserva los datos.",
    )
    for text in tests:
        sf.add_step(doc, text, steps_id)

    sf.add_table(
        doc,
        ["Prueba", "Resultado esperado", "Hecho"],
        [
            ("Caso Bogotá", "4 paneles y techo viable", "[ ]"),
            ("Más consumo", "La necesidad no baja", "[ ]"),
            ("Techo pequeño", "Aparece advertencia de restricción", "[ ]"),
            ("Criterio mensual", "La recomendación no es menor que la anual", "[ ]"),
            ("PVGIS y PVWatts", "Se muestran ambos resultados", "[ ]"),
            ("Exportar e importar", "Se conservan los datos principales", "[ ]"),
        ],
        [2300, 5900, 1160],
        font_size=9.3,
    )
    add_note(
        doc,
        "Si algo falla,",
        "anote el navegador, la fecha, lo que hizo y el mensaje exacto. Una captura también ayuda. No cambie el resultado para que parezca correcto.",
    )


def add_faq_and_script(doc: Document) -> None:
    add_heading(doc, "Preguntas que probablemente tendremos", page_break=True)
    questions = (
        ("Por qué la cobertura puede ser mayor que 100 %", "Porque los paneles son enteros. Al pasar de 3,30 a 4 paneles queda una capacidad adicional."),
        ("Por qué junio puede quedar por debajo de 100 %", "El caso usa balance anual. Los meses con más producción compensan el total del año."),
        ("Por qué el techo puede decir que no es viable", "La recomendación energética puede necesitar más paneles de los que caben en el área escrita."),
        ("La aplicación calcula cuánto dinero ahorraremos", "No. Calcula energía. Para estimar dinero harían falta tarifas, autoconsumo, excedentes y reglas vigentes."),
        ("La aplicación confirma que podemos instalar", "No. Un profesional debe revisar estructura, sombras, electricidad, protecciones y normativa."),
        ("Necesitamos Supabase", "No para esta versión. Los datos se guardan en el navegador y se pueden exportar como archivos."),
    )
    for question, answer in questions:
        add_heading(doc, question, level=2)
        sf.add_para(doc, answer)

    add_heading(doc, "Cómo explicarlo en un minuto", level=2)
    sf.add_para(
        doc,
        "Construimos SolarFit Hogar, una aplicación que estima cuántos paneles solares necesita un hogar a partir de su consumo, "
        "la energía solar disponible, la potencia del panel y el área del techo. Primero calcula una cantidad continua, después redondea "
        "a paneles completos y verifica si caben. La derivada nos permite explicar cómo cambia la necesidad cuando cambia el rendimiento solar. "
        "Probamos el sistema con cinco simulaciones calibradas y comparamos PVGIS con PVWatts. Sirve como una estimación inicial, "
        "no como un diseño final de instalación.",
    )

    add_heading(doc, "Fuentes que usamos", level=2)
    sf.add_para(
        doc,
        "UPME y CORPOEMA para los consumos de referencia; la ficha oficial LONGi para el panel; PVGIS como estimación solar principal; "
        "y PVWatts como segunda estimación. Los detalles completos están guardados en la evidencia del proyecto.",
    )
    add_note(
        doc,
        "Para empezar ahora,",
        f"abra {APP_URL}, cargue Bogotá y después cree un ejercicio cambiando una sola variable.",
    )


def build_document() -> Document:
    cases = load_cases()
    bogota = find_case(cases, "bogota_")

    doc = Document()
    configure_friendly_styles(doc)
    configure_section(doc.sections[0], first_page=True)
    doc.core_properties.title = "Guía sencilla para usar SolarFit Hogar"
    doc.core_properties.subject = "Guía interna para comprender, usar y probar SolarFit Hogar"
    doc.core_properties.author = "Equipo SolarFit Hogar"
    doc.core_properties.keywords = "SolarFit, guía interna, ejemplos, pruebas, paneles solares"

    bullet_id = sf.add_numbering_definition(doc, ordered=False)
    quick_steps_id = sf.add_numbering_definition(doc, ordered=True)
    bogota_steps_id = sf.add_numbering_definition(doc, ordered=True)
    own_steps_id = sf.add_numbering_definition(doc, ordered=True)
    test_steps_id = sf.add_numbering_definition(doc, ordered=True)

    add_cover(doc)
    content_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(content_section)

    add_quick_start(doc, quick_steps_id)
    add_what_we_built(doc, bullet_id)
    add_guided_example(doc, bogota, bogota_steps_id)
    add_own_case(doc, own_steps_id)
    add_practice_ideas(doc, bullet_id)
    add_results_explanation(doc)
    add_test_plan(doc, test_steps_id)
    add_faq_and_script(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return doc


if __name__ == "__main__":
    document = build_document()
    print(f"Guía creada: {OUTPUT}")
    print(f"Secciones: {len(document.sections)} | Párrafos: {len(document.paragraphs)}")
