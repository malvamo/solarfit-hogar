from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = PROJECT_ROOT.parents[2]
PROJECT = WORKSPACE / "Proyectos" / "CalculoDiferencial" / "solarfit-hogar"
OUTPUT = WORKSPACE / "Entregables" / "CalculoDiferencial" / "Proyecto_Completo_SolarFit_Hogar.docx"
RESULTS_IMAGE = PROJECT / "docs" / "assets" / "solarfit-resultados.png"

# narrative_proposal tokens, resolved from the Documents skill preset.
PAGE_W = Inches(8.5)
PAGE_H = Inches(11)
MARGIN = Inches(1)
HEADER_DISTANCE = Inches(0.492)
FOOTER_DISTANCE = Inches(0.492)
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_START_END_DXA = 120

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "22313B"
MUTED = "5C6770"
LIGHT_FILL = "F4F6F9"
MID_FILL = "E8EEF5"
TEAL = "174C4F"
PALE_TEAL = "E6F1EE"
GOLD = "B57A16"
PALE_GOLD = "FFF4D6"
RED = "9B1C1C"
PALE_RED = "FFF1F0"
WHITE = "FFFFFF"
BORDER = "B9C5CF"


def fmt_decimal(value: float, decimals: int = 2) -> str:
    """Format a number for Spanish-language academic tables."""
    text = f"{value:,.{decimals}f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")

# Casos reproducibles ejecutados con la versión 2026.2 del escenario. El consumo
# no representa cinco viviendas medidas: se calibra con promedios oficiales de
# la UPME/CORPOEMA y se mantiene constante durante los doce meses porque la
# fuente no publica una estacionalidad mensual que permita variarlo sin inventar.
CASES = (
    {
        "case": "Bogotá · piso frío",
        "demand_basis": "SIN, > 1.500 m s. n. m.: 156 kWh/mes",
        "annual_demand": 1872.0,
        "pvgis_yield": 1260.65,
        "pvwatts_yield": 1271.06,
        "continuous_panels": 3.299885,
        "pvgis_panels": 4,
        "pvwatts_panels": 4,
        "conservative_panels": 5,
        "roof_area": 20.0,
        "roof_max": 9,
        "buildable_panels": 4,
        "production": 2269.17,
        "coverage": 121.216,
        "minimum_coverage": 98.17,
        "critical_month": "junio",
    },
    {
        "case": "Barranquilla · piso cálido",
        "demand_basis": "SIN, < 500 m s. n. m.: 279 kWh/mes",
        "annual_demand": 3348.0,
        "pvgis_yield": 1567.18,
        "pvwatts_yield": 1643.33,
        "continuous_panels": 4.747381,
        "pvgis_panels": 5,
        "pvwatts_panels": 5,
        "conservative_panels": 6,
        "roof_area": 35.0,
        "roof_max": 16,
        "buildable_panels": 5,
        "production": 3526.16,
        "coverage": 105.321,
        "minimum_coverage": 92.31,
        "critical_month": "noviembre",
    },
    {
        "case": "Medellín · piso templado",
        "demand_basis": "SIN, 500-1.500 m s. n. m.: 178 kWh/mes",
        "annual_demand": 2136.0,
        "pvgis_yield": 1387.56,
        "pvwatts_yield": 1369.68,
        "continuous_panels": 3.420873,
        "pvgis_panels": 4,
        "pvwatts_panels": 4,
        "conservative_panels": 4,
        "roof_area": 14.0,
        "roof_max": 6,
        "buildable_panels": 4,
        "production": 2497.61,
        "coverage": 116.929,
        "minimum_coverage": 100.12,
        "critical_month": "noviembre",
    },
    {
        "case": "Cali · piso templado",
        "demand_basis": "SIN, 500-1.500 m s. n. m.: 178 kWh/mes",
        "annual_demand": 2136.0,
        "pvgis_yield": 1235.28,
        "pvwatts_yield": 1422.18,
        "continuous_panels": 3.842584,
        "pvgis_panels": 4,
        "pvwatts_panels": 4,
        "conservative_panels": 5,
        "roof_area": 18.0,
        "roof_max": 8,
        "buildable_panels": 4,
        "production": 2223.50,
        "coverage": 104.097,
        "minimum_coverage": 90.90,
        "critical_month": "junio",
    },
    {
        "case": "Quibdó rural · cálido y techo restringido",
        "demand_basis": "Rural cálido: 118,8 kWh/mes",
        "annual_demand": 1425.6,
        "pvgis_yield": 898.28,
        "pvwatts_yield": 1131.87,
        "continuous_panels": 3.526740,
        "pvgis_panels": 4,
        "pvwatts_panels": 3,
        "conservative_panels": 5,
        "roof_area": 8.0,
        "roof_max": 3,
        "buildable_panels": 3,
        "production": 1212.68,
        "coverage": 85.064,
        "minimum_coverage": 74.05,
        "critical_month": "junio",
    },
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (
        ("top", CELL_TOP_BOTTOM_DXA),
        ("bottom", CELL_TOP_BOTTOM_DXA),
        ("start", CELL_START_END_DXA),
        ("end", CELL_START_END_DXA),
    ):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = BORDER, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def set_table_geometry(table, widths_dxa: Sequence[int]) -> None:
    if sum(widths_dxa) != CONTENT_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_DXA}: {widths_dxa}")
    if any(len(row.cells) != len(widths_dxa) for row in table.rows):
        raise ValueError("Column count does not match geometry")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)

    set_table_borders(table)


def add_numbering_definition(doc: Document, *, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    existing_abs = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(existing_abs, default=0) + 1
    existing_num = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{0xA0B00000 + abstract_id:08X}")
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    tmpl = OxmlElement("w:tmpl")
    tmpl.set(qn("w:val"), f"{0xC0D00000 + abstract_id:08X}")
    abstract.append(tmpl)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def set_run_font(run, *, name: str = "Calibri", size: float | None = None,
                 color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border(paragraph, color: str, *, left: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    if left:
        node = OxmlElement("w:left")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "18")
        node.set(qn("w:space"), "8")
        node.set(qn("w:color"), color)
        p_bdr.append(node)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False

    for style_name, size, color, bold, align in (
        ("SolarFit Kicker", 9.5, GOLD, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("SolarFit Cover Title", 31, TEAL, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("SolarFit Cover Subtitle", 15, DARK_BLUE, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("SolarFit Lead", 12.5, TEAL, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("SolarFit Equation", 11, DARK_BLUE, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("SolarFit Editable", 10.5, RED, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("SolarFit Source", 9, MUTED, False, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        if style_name not in styles:
            styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[style_name]
        style.font.name = "Cambria Math" if style_name == "SolarFit Equation" else "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.paragraph_format.alignment = align
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.15 if style_name != "SolarFit Lead" else 1.25


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def set_page_geometry(section) -> None:
    section.page_width = PAGE_W
    section.page_height = PAGE_H
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE


def configure_cover_section(section) -> None:
    set_page_geometry(section)
    section.different_first_page_header_footer = False
    section.header.paragraphs[0].text = ""
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Documento académico editable - no sustituye un diseño técnico certificado")
    set_run_font(r, size=8.5, color=MUTED, italic=True)


def configure_content_section(section) -> None:
    set_page_geometry(section)
    section.different_first_page_header_footer = False

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    r = p.add_run("SOLARFIT HOGAR")
    set_run_font(r, size=8.5, color=TEAL, bold=True)
    r = p.add_run("\tCÁLCULO DIFERENCIAL - 2026-2S")
    set_run_font(r, size=8.5, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.15))
    r = p.add_run("Informe acumulativo CAP+IA")
    set_run_font(r, size=8.5, color=MUTED)
    r = p.add_run("\t")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_field(p)

def add_para(doc: Document, text: str, *, bold_lead: str | None = None,
             italic: bool = False, align=None, before: float | None = None,
             after: float | None = None, keep_with_next: bool = False) -> object:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        tail = p.add_run(text[len(bold_lead):])
        set_run_font(tail, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, italic=italic)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep_with_next
    return p


def add_bullet(doc: Document, text: str, bullet_id: int) -> object:
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    marker = p.add_run("•\t")
    set_run_font(marker, color=TEAL, bold=True)
    set_run_font(p.add_run(text))
    return p


def add_step(doc: Document, text: str, decimal_id: int) -> object:
    counters = getattr(doc, "_solarfit_step_counters", {})
    number = counters.get(decimal_id, 0) + 1
    counters[decimal_id] = number
    doc._solarfit_step_counters = counters

    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.36)
    p.paragraph_format.first_line_indent = Inches(-0.26)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    marker = p.add_run(f"{number}.\t")
    set_run_font(marker, color=TEAL, bold=True)
    set_run_font(p.add_run(text))
    return p


def add_callout(doc: Document, label: str, text: str, *, fill: str = PALE_TEAL,
                accent: str = TEAL) -> object:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.keep_together = True
    set_paragraph_shading(p, fill)
    set_paragraph_border(p, accent)
    r = p.add_run(f"{label}. ")
    set_run_font(r, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, color=INK)
    return p


def add_editable(doc: Document, text: str) -> object:
    p = doc.add_paragraph(style="SolarFit Editable")
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.15
    set_paragraph_shading(p, PALE_RED)
    set_paragraph_border(p, RED)
    p.add_run(f"CAMPO EDITABLE DEL GRUPO: {text}")
    return p


def add_equation(doc: Document, text: str, *, note: str | None = None) -> object:
    p = doc.add_paragraph(style="SolarFit Equation")
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4 if note else 9)
    p.paragraph_format.keep_together = True
    set_paragraph_shading(p, LIGHT_FILL)
    p.add_run(text)
    if note:
        n = doc.add_paragraph(note, style="SolarFit Source")
        n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        n.paragraph_format.space_before = Pt(0)
        n.paragraph_format.space_after = Pt(8)
    return p


def add_table(doc: Document, headers: Sequence[str], rows: Iterable[Sequence[str]],
              widths: Sequence[int], *, font_size: float = 9.5) -> object:
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(value)
        set_run_font(r, size=9.2, color=DARK_BLUE, bold=True)
        set_cell_shading(cell, LIGHT_FILL)
    mark_header_row(table.rows[0])
    prevent_row_split(table.rows[0])

    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for cell, value in zip(row.cells, row_values):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)
    return table


def add_table_source(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="SolarFit Source")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)


def add_figure(doc: Document, image_path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Inches(6.12))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", "Pantalla de resultados del prototipo SolarFit Hogar con procedencia, resultados energéticos y advertencias del escenario seleccionado.")
    c = doc.add_paragraph(caption, style="Caption")
    c.paragraph_format.keep_with_next = False


def add_section_title(doc: Document, text: str, *, page_break: bool = False) -> None:
    p = doc.add_paragraph(text, style="Heading 1")
    if page_break:
        p.paragraph_format.page_break_before = True


def add_cover(doc: Document) -> None:
    for _ in range(5):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph("PROYECTO INTEGRADOR CAP+IA", style="SolarFit Kicker")
    p.paragraph_format.space_after = Pt(16)

    p = doc.add_paragraph("SolarFit Hogar", style="SolarFit Cover Title")
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph(
        "Dimensionamiento preliminar de sistemas fotovoltaicos residenciales\nmediante cálculo diferencial",
        style="SolarFit Cover Subtitle",
    )
    p.paragraph_format.space_after = Pt(22)

    p = doc.add_paragraph("INFORME FINAL ACUMULATIVO - HITOS 1, 2 Y 3", style="SolarFit Kicker")
    p.paragraph_format.space_after = Pt(46)

    for text, size, color, bold in (
        ("Cálculo Diferencial - Código 502116 - Grupo 12", 11.5, TEAL, True),
        ("Fundación Universidad de Bogotá Jorge Tadeo Lozano", 11, INK, False),
        ("Docente: Orlando Aya Corredor", 10.5, MUTED, False),
        ("Periodo académico: 2026-2S", 10.5, MUTED, False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(text), size=size, color=color, bold=bold)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("Integrantes: [NOMBRES COMPLETOS Y CÓDIGOS]"), size=10.5, color=RED, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    set_run_font(p.add_run("Versión corregida preparada el 7 de septiembre de 2026"), size=9.5, color=MUTED, italic=True)
def build_document() -> Document:
    doc = Document()
    configure_styles(doc)
    doc.settings.odd_and_even_pages_header_footer = False
    configure_cover_section(doc.sections[0])
    doc.core_properties.title = "SolarFit Hogar - Informe final acumulativo"
    doc.core_properties.subject = "Proyecto integrador CAP+IA de Cálculo Diferencial"
    doc.core_properties.author = "Equipo SolarFit Hogar"
    doc.core_properties.keywords = "cálculo diferencial, energía solar, modelación, CAP+IA"
    doc.core_properties.comments = "Documento académico acumulativo con campos pendientes claramente identificados."

    conventions_id = add_numbering_definition(doc, ordered=True)
    objectives_id = add_numbering_definition(doc, ordered=True)
    checkpoint_id = add_numbering_definition(doc, ordered=True)
    algorithm_id = add_numbering_definition(doc, ordered=True)
    limitations_id = add_numbering_definition(doc, ordered=True)
    conclusions_id = add_numbering_definition(doc, ordered=True)
    poster_id = add_numbering_definition(doc, ordered=True)
    references_id = add_numbering_definition(doc, ordered=True)
    protocol_id = add_numbering_definition(doc, ordered=True)

    add_cover(doc)
    content_section = doc.add_section(WD_SECTION.NEW_PAGE)
    content_section.header.is_linked_to_previous = False
    content_section.footer.is_linked_to_previous = False
    configure_content_section(content_section)

    add_section_title(doc, "Ficha de control y estado del documento")
    add_callout(
        doc,
        "Lectura obligatoria",
        "Este archivo integra el planteamiento, el desarrollo matemático esencial, el prototipo y cinco simulaciones calibradas reproducibles. La demanda se basa en promedios oficiales UPME/CORPOEMA; la producción se estima con PVGIS y se contrasta con PVWatts. Los resultados no son mediciones de cinco viviendas ni una validación empírica: son aproximaciones documentadas para probar el sistema sin inventar datos domésticos.",
        fill=PALE_GOLD,
        accent=GOLD,
    )
    add_table(
        doc,
        ["Elemento", "Estado verificable"],
        [
            ("Planteamiento y modelo", "Desarrollado y revisado en esta versión acumulativa."),
            ("Prototipo web", "Implementado en Next.js/TypeScript; 38 pruebas y compilación productiva aprobadas el 07/09/2026."),
            ("Demanda de los cinco casos", "Calibrada con promedios SUI 2018 publicados por UPME/CORPOEMA; no son facturas ni promedios municipales."),
            ("Panel de referencia", "LONGi LR5-54HTH de 450 W; potencia y dimensiones tomadas de ficha oficial."),
            ("Estimación solar", "PVGIS ejecutado en cinco puntos de referencia, con 14 % de pérdidas y configuración conservada en la evidencia."),
            ("Contraste independiente", "PVWatts ejecutado con la misma capacidad instalable, inclinación, azimut y pérdidas."),
            ("Autoría del grupo", "Pendiente: registrar nombres, aportes reales, decisiones y revisión humana."),
        ],
        [2700, 6660],
    )
    add_editable(doc, "reemplazar esta versión por la fecha real de entrega y confirmar en AVATA el plazo vigente.")
    add_editable(doc, "completar nombres, códigos y aportes verificables de cada integrante antes de entregar.")

    add_section_title(doc, "Resumen ejecutivo", page_break=True)
    add_para(
        doc,
        "SolarFit Hogar responde una pregunta concreta: ¿cuántos paneles solares necesita una vivienda conectada a la red para compensar una proporción de su consumo eléctrico? La respuesta se formula como una prefactibilidad energética y no como un diseño de instalación. El proyecto combina doce consumos mensuales, recurso solar, potencia y dimensiones de un módulo, meta de cobertura y área útil del techo. Conserva la solución continua para analizarla con cálculo diferencial y solo al final la transforma en una cantidad entera de paneles.",
    )
    add_para(
        doc,
        "El modelo ofrece dos rutas mutuamente excluyentes. La primera calcula la producción a partir de horas solares pico, días y un factor global de rendimiento. La segunda usa el rendimiento específico neto entregado por PVGIS; en esta ruta no se vuelve a aplicar el factor de rendimiento porque se duplicarían las pérdidas. El usuario puede seleccionar balance anual o un criterio conservador que atiende el mes más exigente. La recomendación energética se contrasta con la capacidad física del techo sin ocultar una posible inviabilidad.",
    )
    add_para(
        doc,
        "La implementación es una aplicación web modular con formularios, importación CSV y JSON, visualizaciones, sensibilidad, incertidumbre, exportaciones y persistencia local. El 7 de septiembre de 2026 se ejecutó la comprobación completa del repositorio: análisis estático, verificación de TypeScript, 38 pruebas en 9 archivos y compilación productiva; todas finalizaron correctamente. Esta evidencia prueba coherencia interna del software, no exactitud empírica para una vivienda concreta.",
    )
    add_callout(
        doc,
        "Aporte central",
        "La aplicación no entrega únicamente un número. Expone variables, unidades, procedencia, supuestos, solución continua, redondeo, cobertura mensual, sensibilidad, incertidumbre y restricción de área para que el resultado pueda explicarse y auditarse.",
    )
    add_para(
        doc,
        "El informe se organiza de forma acumulativa. El Hito 1 presenta el problema y el primer modelo; el Hito 2 desarrolla únicamente la derivada que cambia una decisión del proyecto, además del algoritmo y las visualizaciones; el Hito 3 documenta el producto, las cinco simulaciones calibradas, el contraste entre fuentes y los anexos de cierre.",
    )

    add_section_title(doc, "Guía de lectura y trazabilidad", page_break=True)
    add_para(doc, "La estructura reproduce la evolución exigida por el proyecto CAP+IA y permite distinguir contenido terminado de evidencia pendiente.")
    add_table(
        doc,
        ["Bloque", "Contenido", "Relación con la evaluación"],
        [
            ("Hito 1", "Contexto, pregunta, objetivos, fuentes, modelo inicial, organización y riesgos.", "Formulación y diseño."),
            ("Hito 2", "Desarrollo matemático, resultados intermedios, algoritmo, arquitectura, visualización y ajustes.", "Investigación y creación."),
            ("Hito 3", "Producto, casos calibrados, contraste, reflexión, conclusiones, póster y pitch.", "Comunicación y evaluación."),
            ("Anexos", "Plantilla para un caso propio, trazabilidad, pruebas, IA y contribuciones.", "Soporte verificable."),
        ],
        [1600, 4560, 3200],
    )
    add_para(doc, "Convenciones de estado:", bold_lead="Convenciones de estado:")
    add_step(doc, "Implementado: existe en el prototipo y fue comprobado mediante el comando de verificación del repositorio.", conventions_id)
    add_step(doc, "Simulación calibrada: dato construido con un promedio oficial y supuestos explícitos; no representa una vivienda medida.", conventions_id)
    add_step(doc, "Pendiente del grupo: requiere medición, decisión o autoría que no puede inventarse.", conventions_id)
    add_callout(
        doc,
        "Conflicto documental",
        "El sílabo ubica el Hito 1 en la semana 4 y el Hito 2 en la semana 9, mientras la rúbrica general los denomina Momento 1 - Semana 5 y Momento 2 - Semana 10. La tercera entrega sí coincide en semana 14. El grupo debe confirmar en AVATA la fecha operativa y no elegir silenciosamente una versión.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    # HITO 1
    add_section_title(doc, "Hito 1 - Diseño y desafío", page_break=True)
    add_para(doc, "Esta sección consolida y corrige el primer entregable. Mantiene la idea aprobada por el equipo y el profesor, y refuerza la conexión entre la pregunta de energía solar, la modelación y la Ingeniería de Sistemas.")

    doc.add_paragraph("1. Contextualización del problema", style="Heading 2")
    add_para(
        doc,
        "Dimensionar preliminarmente un sistema fotovoltaico residencial no consiste en dividir el consumo por la potencia nominal del panel. La potencia se expresa en kW bajo condiciones de ensayo, mientras la factura registra energía en kWh durante un periodo. La energía que puede producir un módulo depende del recurso solar, los días, la orientación, la inclinación, las pérdidas y las restricciones físicas. Confundir potencia con energía o usar horas de luz como si fueran horas solares pico conduce a resultados engañosos.",
    )
    add_para(
        doc,
        "Desde Ingeniería de Sistemas, el problema exige administrar datos heterogéneos, validar unidades y rangos, formular funciones, convertirlas en un algoritmo, separar capas de software, visualizar resultados y conservar trazabilidad. SolarFit Hogar se concibe como apoyo a decisiones: reduce la fricción de una evaluación inicial y deja visibles los supuestos para que el usuario entienda la recomendación.",
    )
    add_callout(doc, "Pregunta de investigación", "Para una vivienda conectada a la red en Colombia, ¿cuántos paneles de una potencia dada se requieren para compensar una proporción objetivo del consumo, según el rendimiento solar específico y el área útil del techo?")

    doc.add_paragraph("1.1 Usuarios y propuesta de valor", style="Heading 3")
    add_table(
        doc,
        ["Actor", "Necesidad", "Valor preliminar"],
        [
            ("Propietario", "Comprender el orden de magnitud antes de solicitar cotizaciones.", "Escenarios explicables y advertencias."),
            ("Instalador", "Estandarizar el levantamiento inicial.", "Reporte trazable para una conversación comercial."),
            ("Equipo académico", "Aplicar cálculo diferencial a un fenómeno contextualizado.", "Función n(Y), primera derivada, algoritmo, gráfica y contraste."),
        ],
        [1800, 3600, 3960],
    )
    add_para(doc, "La oportunidad de startup es una hipótesis, no un hallazgo validado. Una evolución comercial podría ofrecer evaluaciones iniciales gratuitas y herramientas profesionales para instaladores, pero requeriría entrevistas, análisis normativo, estructura de costos y validación de mercado.")

    doc.add_paragraph("2. Objetivos", style="Heading 2")
    add_para(doc, "Objetivo general. Desarrollar y contrastar un modelo de prefactibilidad que estime la capacidad fotovoltaica y el número de paneles necesarios para compensar una proporción del consumo eléctrico de una vivienda conectada a la red, incorporando el rendimiento solar y el área disponible.", bold_lead="Objetivo general.")
    add_para(doc, "Objetivos específicos:", bold_lead="Objetivos específicos:")
    for text in (
        "Caracterizar el consumo mediante facturas trazables o promedios oficiales claramente rotulados como simulación calibrada.",
        "Obtener datos solares trazables y especificaciones oficiales de al menos un módulo comercial.",
        "Formular energía, cobertura y cantidad continua de paneles con variables, unidades, dominio y restricciones.",
        "Interpretar la primera derivada de n(Y) para cuantificar cómo cambia la necesidad continua de paneles cuando cambia el rendimiento solar.",
        "Implementar un prototipo web que compare criterios, redondee al final y verifique el área del techo.",
        "Contrastar PVGIS con PVWatts bajo parámetros comparables y documentar las diferencias y limitaciones.",
    ):
        add_step(doc, text, objectives_id)

    doc.add_paragraph("3. Alcance, exclusiones y supuestos", style="Heading 2")
    add_table(
        doc,
        ["Incluido", "Excluido"],
        [
            ("Sistema residencial conectado a red y sin baterías.", "Autonomía nocturna o respaldo de cargas."),
            ("Balance energético mensual y anual.", "Coincidencia horaria entre producción y consumo."),
            ("Potencia, número de paneles y restricción de área.", "Inversor, cadenas, protecciones, cables y puesta a tierra."),
            ("Inclinación y orientación parametrizables.", "Inspección de sombras cercanas y cálculo estructural."),
            ("Prefactibilidad y escenarios.", "Certificación, cotización definitiva o promesa de factura cero."),
        ],
        [4680, 4680],
    )
    add_callout(doc, "Límite profesional", "La recomendación energética no autoriza una instalación. Un diseño real requiere visita técnica y revisión eléctrica, estructural, normativa y de seguridad por profesionales competentes.", fill=PALE_GOLD, accent=GOLD)

    sources_heading = doc.add_paragraph("4. Fuentes y plan de datos", style="Heading 2")
    sources_heading.paragraph_format.page_break_before = True
    add_para(doc, "El modelo separa cada magnitud de su evidencia. En los cinco casos de prueba, el consumo mensual se fija con promedios oficiales UPME/CORPOEMA; se repite el promedio en los doce meses porque esa fuente no publica una estacionalidad mensual. Esto evita fabricar variaciones. El grupo conserva además la ficha del panel y las respuestas de PVGIS y PVWatts.", keep_with_next=True)
    add_table(
        doc,
        ["Dato", "Fuente prevista", "Control mínimo", "Estado"],
        [
            ("Consumo", "UPME/CORPOEMA, promedios SUI 2018", "Piso térmico/sector y página de la tabla", "Calibrado"),
            ("Ubicación", "Puntos urbanos de referencia", "Supuesto reproducible; no son viviendas", "Documentado"),
            ("Recurso solar", "PVGIS 5.3", "Inclinación 10°, azimut 180° y pérdidas 14 %", "Ejecutado"),
            ("Panel", "Ficha LONGi LR5-54HTH", "450 W; 1,722 m × 1,134 m", "Oficial"),
            ("Techo", "Áreas de escenario", "Supuesto de prueba; no son mediciones", "Documentado"),
            ("Contraste", "PVWatts V8", "Misma capacidad instalable y parámetros", "Ejecutado"),
        ],
        [1500, 2800, 3260, 1800],
        font_size=8.8,
    )
    add_callout(doc, "Alcance de los datos", "Los promedios UPME/CORPOEMA no son consumos de hogares concretos ni promedios de cada ciudad. Se usan para calibrar casos plausibles por piso térmico y sector. Las áreas de techo y los puntos de referencia son supuestos de prueba, no mediciones.", fill=PALE_GOLD, accent=GOLD)

    doc.add_paragraph("5. Primer modelo matemático", style="Heading 2")
    add_para(doc, "Se define el índice mensual m = 1, ..., 12; Cₘ es el consumo mensual, P la potencia del panel en kWp, q la meta de cobertura como fracción, dₘ los días y Aᵤ el área útil. La energía de un panel se representa por Gₘ y la cantidad continua requerida por nₘ*. Los dos modos solares no deben mezclarse. Para el análisis diferencial se usa el rendimiento específico anual Y y la demanda anual D.")
    add_equation(doc, "Modo manual:     Gₘ = P · hₘ · dₘ · PRₘ")
    add_equation(doc, "Modo PVGIS:     Gₘ = P · Yₘ⁽ⁿᵉᵗ⁾", note="Yₘ⁽ⁿᵉᵗ⁾ se expresa en kWh/kWp y ya incorpora las pérdidas configuradas en la simulación.")
    add_equation(doc, "nₘ* = q · Cₘ / Gₘ")
    add_equation(doc, "n*ₐₙᵤₐₗ = q · ΣCₘ / ΣGₘ        |        n*conservador = máx. (nₘ*)")
    add_equation(doc, "N = ⌈n*⌉        |        Aₚₐₙₑₗ,inst = ancho · alto · fₐᵣₑₐ        |        Nₘₐₓ = ⌊Aᵤ / Aₚₐₙₑₗ,inst⌋")
    add_equation(doc, "n(Y) = q · D / (P · Y)", note="Y es el rendimiento específico neto anual en kWh/kWp-año. Esta es la única función usada para el análisis diferencial.")
    add_callout(doc, "Decisión correcta", "La derivada se calcula sobre n(Y), antes del redondeo. La función ceil se aplica una sola vez al final. Si el techo admite menos paneles, el sistema conserva la recomendación matemática y reporta por separado la cantidad instalable.")

    doc.add_paragraph("6. Organización, riesgos y criterios de éxito", style="Heading 2")
    add_table(
        doc,
        ["Riesgo", "Efecto", "Respuesta"],
        [
            ("Serie incompleta o sin regla", "Baja representatividad", "Usar datos disponibles o simulación explícita; nunca presentar supuestos como medición."),
            ("HSP mal interpretada", "Producción sobreestimada", "Verificar unidad y fuente; no usar horas de luz."),
            ("Pérdidas duplicadas", "Más paneles de los necesarios", "Aplicar PR solo en modo manual."),
            ("Techo impreciso", "Falsa viabilidad", "Medir área útil y documentar obstáculos."),
            ("Dependencia de IA", "Errores no detectados", "Pruebas, cálculo manual y revisión del equipo."),
        ],
        [2000, 2800, 4560],
        font_size=9.0,
    )
    add_editable(doc, "asignar responsables reales de coordinación, datos, modelación, programación, visualización, redacción y pitch.")
    doc.add_paragraph("Punto de control del Hito 1", style="Heading 3")
    for text in (
        "Problema delimitado y conectado con Ingeniería de Sistemas.",
        "Objetivos medibles y modelo inicial con unidades y restricciones.",
        "Fuentes previstas, riesgos y campos de responsabilidad identificados.",
        "Cinco casos calibrados definidos con datos oficiales, supuestos trazables y alcance explícito.",
    ):
        add_step(doc, text, checkpoint_id)

    # HITO 2
    add_section_title(doc, "Hito 2 - Investigación y creación", page_break=True)
    add_para(doc, "El segundo hito transforma el modelo inicial en un procedimiento reproducible. De acuerdo con la retroalimentación del profesor, el análisis diferencial se limita a la función y la primera derivada que sí cambian la interpretación del dimensionamiento. Los demás cálculos son balances energéticos, restricciones o comparaciones de apoyo.")

    doc.add_paragraph("7. Desarrollo matemático esencial", style="Heading 2")
    doc.add_paragraph("7.1 Función directamente asociada con la decisión", style="Heading 3")
    add_para(doc, "Sea D la demanda de energía del periodo, q la fracción de consumo que se desea compensar, P la potencia nominal de un panel y Y el rendimiento específico neto del sitio para ese mismo periodo. La cantidad continua de paneles es:")
    add_equation(doc, "n(Y) = q · D / (P · Y),     con Y > 0")
    add_para(doc, "Esta función conecta una entrada solar que cambia según el sitio y la configuración con la salida principal del proyecto. Primero se interpreta n(Y) como capacidad continua; después se aplica N = ⌈n(Y)⌉ para obtener paneles enteros. Así, el redondeo no distorsiona el análisis.")

    doc.add_paragraph("7.2 Primera derivada e interpretación práctica", style="Heading 3")
    add_equation(doc, "n'(Y) = -q · D / (P · Y²) < 0")
    add_para(doc, "La derivada mide el cambio local de la necesidad continua de paneles por cada unidad adicional de rendimiento específico. Su signo negativo tiene impacto directo: si PVGIS estima un mayor rendimiento, se requieren menos paneles continuos para la misma demanda, potencia y meta. La aplicación muestra esta sensibilidad antes del redondeo para que el usuario entienda si la recomendación está cerca de cambiar de entero.")
    add_callout(doc, "Ejemplo con Bogotá", "Con D = 1.872 kWh/año, P = 0,45 kWp, q = 1 y Y = 1.260,65 kWh/kWp-año, se obtiene n(Y) = 3,300 y n'(Y) ≈ -0,002618 paneles por cada kWh/kWp-año. Cerca de ese punto, 100 unidades adicionales de rendimiento reducen la necesidad continua en aproximadamente 0,262 paneles. La decisión entera se recalcula después con ceil.")
    add_callout(doc, "Alcance del cálculo diferencial", "El informe conserva únicamente la función n(Y) y su primera derivada porque ambas explican una variación con impacto directo en la decisión. Los contenidos que no cambian el dimensionamiento quedan fuera, de acuerdo con la retroalimentación docente.", fill=PALE_GOLD, accent=GOLD)

    doc.add_paragraph("7.3 Comparación discreta de inclinaciones", style="Heading 3")
    add_para(doc, "La aplicación consulta en PVGIS una lista finita de inclinaciones candidatas y compara su rendimiento anual. Reporta el mejor valor entre los candidatos consultados. Este procedimiento es una comparación discreta de escenarios y no una aplicación adicional de cálculo diferencial.")

    doc.add_paragraph("7.4 Escenario de incertidumbre", style="Heading 3")
    add_para(doc, "El rango configurado de incertidumbre es un apoyo para explorar cómo una variación del rendimiento podría cambiar el número entero recomendado. No se presenta como tema de cálculo diferencial ni como intervalo estadístico de confianza.")

    doc.add_paragraph("8. Variables, unidades y controles", style="Heading 2")
    add_table(
        doc,
        ["Símbolo", "Significado", "Unidad/dominio", "Rol"],
        [
            ("Cₘ", "Consumo del mes m", "kWh/mes; ≥ 0", "Entrada"),
            ("P", "Potencia nominal del panel", "kWp; > 0", "Entrada"),
            ("hₘ", "Horas solares pico diarias", "h/día; > 0", "Modo manual"),
            ("dₘ", "Días del periodo", "día/mes; 1 a 45", "Entrada"),
            ("PRₘ", "Factor global de rendimiento", "0 < PR ≤ 1", "Solo modo manual"),
            ("Yₘ⁽ⁿᵉᵗ⁾", "Rendimiento específico neto", "kWh/kWp-mes; > 0", "Modo PVGIS"),
            ("Y", "Rendimiento específico neto del periodo", "kWh/kWp-periodo; > 0", "Variable de n(Y)"),
            ("q", "Meta de cobertura", "0 < q ≤ 1", "Decisión"),
            ("Aᵤ", "Área útil del techo", "m²; ≥ 0", "Restricción"),
            ("fₐᵣₑₐ", "Factor de separación/montaje", "1 a 3", "Parámetro"),
            ("n*", "Paneles continuos equivalentes", "paneles; ≥ 0", "Salida analítica"),
            ("N", "Paneles enteros recomendados", "entero; ≥ 0", "Decisión"),
        ],
        [1300, 3300, 2660, 2100],
        font_size=8.6,
    )
    add_table_source(doc, "Fuente: definiciones del modelo implementado en src/domain/solar/types.ts y calculate.ts.")

    doc.add_paragraph("9. Algoritmo y reglas de decisión", style="Heading 2")
    for text in (
        "Recibir ubicación general, doce consumos, días, área, meta, panel y configuración solar.",
        "Validar cantidad de meses, unicidad, valores finitos, rangos y unidades.",
        "Seleccionar exactamente un modo solar: HSP manual o rendimiento neto PVGIS.",
        "Calcular el rendimiento mensual por kWp y luego la energía mensual por panel.",
        "Obtener nₘ*, el balance anual y el mes limitante.",
        "Elegir estrategia anual o conservadora; conservar n* y aplicar ceil solo al final.",
        "Calcular Nₘₐₓ del techo y separar recomendación de cantidad instalable.",
        "Generar producción, cobertura, balances, sensibilidad n'(Y), incertidumbre y advertencias.",
        "Comparar inclinaciones candidatas de forma discreta; no inferir un óptimo continuo.",
        "Exportar resultados con procedencia y conservar los soportes fuera de la aplicación.",
    ):
        add_step(doc, text, algorithm_id)
    add_callout(doc, "Caso de consumo cero", "Un mes con Cₘ = 0 produce nₘ* = 0 y una cobertura porcentual indefinida. El sistema conserva el balance energético y representa la cobertura como nula, no como infinito.")

    doc.add_paragraph("10. Arquitectura del prototipo", style="Heading 2")
    add_para(doc, "La aplicación usa Next.js 16, React 19 y TypeScript. El cálculo reside en funciones puras y no en la interfaz, lo que permite pruebas deterministas. La persistencia del borrador ocurre en localStorage; no existe autenticación, base de datos ni telemetría.")
    add_equation(doc, "Entradas validadas → motor matemático puro → resultados + advertencias")
    add_equation(doc, "Navegador → /api/pvgis → PVGIS 5.3        |        Navegador → /api/pvwatts → PVWatts v8")
    add_table(
        doc,
        ["Capa", "Responsabilidad", "Ejemplos"],
        [
            ("Presentación", "Formulario, navegación y estilos", "app/page, globals.css"),
            ("Interacción", "Estado, consulta, importación y exportación", "SolarFitApp"),
            ("Dominio", "Cálculo, contraste, facturas y comparación de inclinaciones", "domain/solar"),
            ("Servicios", "Normalización y proveedores solares", "services/pvgis, pvwatts"),
            ("API", "Control de parámetros y llamadas externas", "api/pvgis, optimize, pvwatts"),
            ("Pruebas", "Casos de cálculo, archivos y proveedores", "9 archivos, 38 pruebas"),
        ],
        [1800, 4300, 3260],
        font_size=9.0,
    )
    add_callout(doc, "Privacidad", "La aplicación pide municipio o coordenadas generales, no dirección. Al consultar servicios solares, las coordenadas y parámetros se transmiten al proveedor externo. El JSON exportado debe revisarse antes de compartirlo.", fill=PALE_GOLD, accent=GOLD)

    doc.add_paragraph("11. Ajustes realizados frente al primer modelo", style="Heading 2")
    add_table(
        doc,
        ["Problema detectado", "Ajuste", "Efecto"],
        [
            ("Posible duplicación de pérdidas", "Dos modos solares excluyentes.", "PVGIS no vuelve a multiplicarse por PR."),
            ("Redondeo prematuro", "Se conserva n*.", "La derivada n'(Y) y la decisión entera son coherentes."),
            ("Techo insuficiente", "Dos cantidades: requerida e instalable.", "La app no oculta la inviabilidad."),
            ("Factura no alineada con mes", "Utilidad de prorrateo diario.", "Permite calendarizar periodos reales."),
            ("Fuente sin trazabilidad", "Metadatos de procedencia.", "Distingue simulación calibrada, usuario y fuente oficial."),
            ("Exceso de temas matemáticos", "Se conserva solo n(Y) y n'(Y).", "Cada concepto de cálculo tiene impacto directo."),
            ("Demanda arbitraria", "Promedios UPME/CORPOEMA y regla constante explícita.", "Casos calibrados sin estacionalidad inventada."),
        ],
        [2500, 3060, 3800],
        font_size=8.9,
    )

    doc.add_paragraph("12. Visualización y resultado intermedio", style="Heading 2")
    add_para(doc, "La interfaz se divide en Datos, Modelo y Resultados. Presenta la recomendación, la capacidad continua, la cobertura anual, la restricción física, una serie mensual y la curva n(Y). El caso inicial está rotulado como simulación calibrada y permite distinguir datos oficiales, estimaciones de proveedor y supuestos de prueba.")
    if RESULTS_IMAGE.exists():
        add_figure(doc, RESULTS_IMAGE, "Figura 1. Pantalla de resultados del prototipo. La captura demuestra la ejecución visual; las cifras deben leerse junto con la procedencia y los supuestos del caso. Elaboración propia.")
    else:
        add_callout(doc, "Figura pendiente", "No se encontró la captura del prototipo en la ruta prevista. Inserte una captura actual de la versión verificada.", fill=PALE_RED, accent=RED)

    # HITO 3
    add_section_title(doc, "Hito 3 - Comunicación y evaluación", page_break=True)
    add_para(doc, "El tercer hito integra el proceso completo, demuestra el producto y documenta cinco simulaciones calibradas. Esta ruta se adopta porque levantar cinco instalaciones reales no es viable en el tiempo del curso. Cada caso conserva la fuente de demanda, los parámetros, las respuestas de PVGIS y PVWatts y las limitaciones de interpretación.")

    doc.add_paragraph("13. Producto final", style="Heading 2")
    add_para(doc, "El producto es un prototipo web de prefactibilidad denominado SolarFit Hogar. Su salida principal es una recomendación explicable de potencia y paneles, acompañada por cantidad continua, producción y cobertura mensuales, balance, sensibilidad, incertidumbre, viabilidad por área y advertencias de alcance.")
    add_table(
        doc,
        ["Capacidad", "Estado al 07/09/2026", "Evidencia"],
        [
            ("Formulario de 12 meses", "Implementado", "Interfaz y dominio"),
            ("Importación CSV", "Implementado", "Parser mensual"),
            ("Cálculo anual y conservador", "Implementado", "calculateSolarProject"),
            ("PVGIS neto", "Implementado", "Ruta /api/pvgis"),
            ("Comparación discreta de inclinaciones", "Implementado", "Interfaz y ruta técnica /api/pvgis/optimize"),
            ("PVWatts independiente", "Implementado", "Interfaz y ruta /api/pvwatts; requiere NLR_API_KEY"),
            ("Gráficos y tabla mensual", "Implementado", "SVG e interfaz"),
            ("Importación/exportación JSON y exportación CSV/PDF", "Implementado", "Restauración, descargas e impresión"),
            ("Cinco simulaciones calibradas", "Ejecutadas", "UPME/CORPOEMA + PVGIS + PVWatts + ficha LONGi"),
        ],
        [2900, 2800, 3660],
        font_size=8.9,
    )
    add_callout(doc, "Resultado de la verificación técnica", "El comando pnpm check finalizó correctamente el 07/09/2026: ESLint sin errores, TypeScript sin errores, 9 archivos de pruebas con 38 pruebas aprobadas y compilación productiva de Next.js completada. Esta verificación demuestra consistencia interna del software, no desempeño medido de una instalación solar.")

    doc.add_paragraph("14. Resultados de los cinco casos calibrados", style="Heading 2")
    add_para(doc, "Los cinco casos usan meta q = 100 %, panel LONGi de 450 W, inclinación de 10°, azimut de 180°, pérdidas de 14 % y factor de área de 1,10. El consumo permanece constante durante doce meses: esa regla reproduce el promedio oficial disponible y evita inventar estacionalidad. Las áreas de techo son supuestos deliberados para probar tanto escenarios viables como uno restringido.")
    add_table(
        doc,
        ["Caso calibrado", "Base mensual de demanda", "Demanda anual", "Techo supuesto"],
        [
            (
                case["case"],
                case["demand_basis"],
                f'{fmt_decimal(case["annual_demand"], 1)} kWh',
                f'{fmt_decimal(case["roof_area"], 1)} m²; máximo {case["roof_max"]} paneles',
            )
            for case in CASES
        ],
        [2100, 3400, 1700, 2160],
        font_size=8.8,
    )
    add_table_source(doc, "Fuente de demanda: UPME/CORPOEMA (2019), promedio SUI 2018 por altura para el SIN (p. 28) y por clima para el sector rural (p. 81). Bogotá, Barranquilla, Medellín y Cali se calibran por rango altitudinal; Quibdó usa el promedio rural cálido. No son promedios municipales ni facturas domésticas.")
    add_table(
        doc,
        ["Caso", "Y PVGIS / PVWatts", "n(Y) y decisión", "Instalable y cobertura", "Mínimo mensual"],
        [
            (
                case["case"],
                f'{fmt_decimal(case["pvgis_yield"])} / {fmt_decimal(case["pvwatts_yield"])} kWh/kWp-año',
                f'{fmt_decimal(case["continuous_panels"], 3)}; N {case["pvgis_panels"]} / {case["pvwatts_panels"]}; conservador {case["conservative_panels"]}',
                f'{case["buildable_panels"]} paneles; {fmt_decimal(case["production"])} kWh; {fmt_decimal(case["coverage"])} %',
                f'{fmt_decimal(case["minimum_coverage"])} % en {case["critical_month"]}',
            )
            for case in CASES
        ],
        [1700, 1900, 2050, 2350, 1360],
        font_size=7.8,
    )
    add_table_source(doc, "Y es rendimiento específico neto anual. La decisión muestra paneles por balance anual con PVGIS / PVWatts; el valor conservador corresponde al mes más exigente. Producción y cobertura usan únicamente la capacidad físicamente instalable en PVGIS.")
    add_callout(doc, "Hallazgo de borde", "Quibdó requiere 4 paneles por balance anual PVGIS, pero el techo supuesto solo admite 3. Por eso la aplicación no promete la meta: informa 85,06 % de cobertura anual estimada y 74,05 % en junio. PVWatts, contrastado sobre esos mismos 3 paneles, entrega un rendimiento mayor y sí cambia la decisión energética a 3; la discrepancia se conserva como incertidumbre entre modelos, no se oculta.", fill=PALE_GOLD, accent=GOLD)
    add_callout(doc, "Interpretación responsable", "Estos resultados son simulaciones calibradas y reproducibles. Sirven para probar el cálculo, comparar escenarios y sustentar el funcionamiento del MVP. No demuestran cuánto producirá una instalación específica ni reemplazan facturas, medición de techo, sombras, inspección estructural o diseño eléctrico.", fill=PALE_GOLD, accent=GOLD)

    doc.add_paragraph("15. Verificación y contraste de resultados", style="Heading 2")
    add_para(doc, "La evidencia se organiza en capas para no confundir corrección del algoritmo con exactitud física. Al no existir mediciones de sistemas instalados, el trabajo no afirma validación empírica.")
    add_table(
        doc,
        ["Capa", "Prueba", "Criterio"],
        [
            ("Dimensional", "Revisar unidades de cada ecuación.", "Gₘ en kWh/panel-mes y n* adimensional en paneles."),
            ("Analítica", "Comprobar n'(Y).", "Al aumentar Y debe disminuir n(Y) antes del redondeo."),
            ("Algorítmica", "Casos manuales y unitarios.", "Coincidencia dentro de tolerancia y redondeo al final."),
            ("Restricciones", "Techo insuficiente y extremos.", "Advertencia explícita; no truncar recomendación."),
            ("Contraste externo", "Comparar PVGIS con PVWatts.", "Misma capacidad instalable, sitio de referencia, inclinación, azimut y pérdidas."),
            ("Trazabilidad", "Reconstruir cada caso.", "Fuente oficial, supuesto, entrada, transformación y salida identificables."),
        ],
        [1600, 3960, 3800],
        font_size=8.8,
    )
    add_para(doc, "El contraste principal usa dos magnitudes comprensibles para la decisión: rendimiento anual estimado por cada proveedor y cantidad entera de paneles resultante. En cuatro casos ambas fuentes producen la misma decisión anual; en Quibdó difieren en un panel. Esta comparación no determina cuál proveedor es empíricamente más exacto porque no hay producción medida en sitio.")

    doc.add_paragraph("16. Limitaciones e interpretación responsable", style="Heading 2")
    for text in (
        "El balance mensual no representa la simultaneidad horaria ni calcula autoconsumo económico.",
        "Cobertura superior a 100 % significa energía anual modelada mayor que el consumo; no equivale a ahorro garantizado ni venta automática a la red.",
        "PVGIS considera horizonte topográfico, pero no sustituye un levantamiento de sombras cercanas.",
        "El factor de área aproxima separaciones; no modela geometría, pasillos, cargas de viento ni estructura.",
        "La ficha de potencia nominal no describe por sí sola temperatura, degradación, tolerancias ni compatibilidad eléctrica.",
        "El criterio conservador puede sobredimensionar si la red permite compensar excedentes entre periodos; su elección debe justificarse.",
        "La comparación de inclinaciones selecciona el mejor candidato probado, no un óptimo continuo.",
        "La incertidumbre configurada es un escenario determinista y no un intervalo probabilístico calibrado.",
        "Los consumos UPME/CORPOEMA son promedios por piso térmico o sector; no describen cada ciudad ni sustituyen las facturas de una vivienda.",
        "Las coordenadas de referencia y áreas de techo son supuestos reproducibles de prueba; no son levantamientos de inmuebles.",
    ):
        add_step(doc, text, limitations_id)

    doc.add_paragraph("17. Evolución corte a corte", style="Heading 2")
    add_table(
        doc,
        ["Corte", "Pregunta dominante", "Avance acumulado", "Evidencia"],
        [
            ("Hito 1", "¿Qué problema se modela?", "Contexto, objetivos, fuentes, variables, primer modelo y plan.", "Anteproyecto revisado."),
            ("Hito 2", "¿Cómo funciona y cambia el modelo?", "n(Y), n'(Y), algoritmo, arquitectura, pruebas y visualizaciones.", "Prototipo verificable."),
            ("Hito 3", "¿Qué evidencia lo respalda y cómo se comunica?", "Casos calibrados, contraste, limitaciones, reflexión, documento, póster y pitch.", "Paquete final trazable."),
        ],
        [1300, 2200, 3660, 2200],
        font_size=8.8,
    )
    add_para(doc, "La coherencia se conserva porque la pregunta original no cambia; cada corte añade evidencia. La corrección más importante fue retirar temas matemáticos que no incidían en la decisión y concentrar el cálculo diferencial en n(Y) y n'(Y). También se distinguieron recomendación energética y capacidad física, y se calibraron los casos con datos oficiales sin presentarlos como mediciones reales.")

    doc.add_paragraph("18. Reflexión crítica sobre matemática e IA", style="Heading 2")
    add_para(doc, "La matemática aporta una estructura explicable. La función n(Y) convierte demanda, potencia y rendimiento en una necesidad continua; su primera derivada cuantifica directamente cómo cambia esa necesidad cuando cambia la estimación solar. El redondeo y la restricción del techo convierten luego el resultado continuo en una decisión física. Esta relación es suficiente para el objetivo del proyecto; agregar otros temas del curso no mejoraría la decisión.")
    add_para(doc, "La IA se utilizó como apoyo para analizar requisitos, proponer alternativas, revisar fórmulas, estructurar código, generar pruebas y preparar documentación. Su utilidad fue acelerar la exploración y detectar riesgos, especialmente la duplicación de pérdidas. Sus respuestas no se aceptaron como autoridad: se contrastaron con los archivos del curso, el código ejecutado, pruebas unitarias y documentación oficial.")
    add_callout(doc, "Aprendizaje", "Una salida computacional puede ser precisa respecto del algoritmo y aun así no describir una vivienda concreta si las entradas son promedios y supuestos. Por eso el informe distingue verificación del código, contraste entre modelos y validación empírica, que aquí no se realizó.")

    doc.add_paragraph("19. Conclusiones", style="Heading 2")
    for text in (
        "La cantidad de paneles depende del consumo, la meta, el recurso solar, la potencia, las pérdidas y el área; no existe una cifra universal para una vivienda.",
        "El modelo manual y el modelo con rendimiento neto son alternativas, no factores acumulables.",
        "La solución continua n(Y) y su primera derivada aportan la sensibilidad directamente útil; ceil convierte esa solución en una decisión física solo al final.",
        "La restricción de techo debe reportarse por separado para no ocultar una meta energéticamente requerida pero físicamente inviable.",
        "El prototipo está técnicamente implementado, verificado internamente y probado con cinco simulaciones calibradas; ninguna equivale a una validación empírica en una vivienda.",
        "El potencial de startup existe como hipótesis de valor y requiere validación comercial, legal y operacional fuera del alcance académico actual.",
    ):
        add_step(doc, text, conclusions_id)

    add_section_title(doc, "Declaraciones obligatorias", page_break=False)
    doc.add_paragraph("20. Declaración de uso de inteligencia artificial", style="Heading 2")
    add_para(doc, "Herramienta utilizada. ChatGPT/Codex de OpenAI.", bold_lead="Herramienta utilizada.")
    add_para(doc, "Propósitos. Análisis de fuentes y rúbricas; organización del trabajo; exploración y revisión del modelo; apoyo en arquitectura, programación, pruebas, redacción y diseño de entregables.", bold_lead="Propósitos.")
    add_para(doc, "Aportes aceptados. Separación de los dos modos solares; conservación de la solución continua; advertencias de techo; diseño modular; pruebas de monotonicidad y de no duplicación de pérdidas; estructura acumulativa del informe.", bold_lead="Aportes aceptados.")
    add_para(doc, "Aportes modificados. Las propuestas iniciales se ajustaron a las unidades del problema, los archivos del curso y la retroalimentación docente: se retiraron cálculos accesorios, se calibró la demanda con una fuente oficial y se distinguió simulación de medición real.", bold_lead="Aportes modificados.")
    add_para(doc, "Aportes rechazados o no incorporados. Cualquier dato numérico no soportado, referencias no verificadas, promesas de ahorro, diseño eléctrico detallado y atribuciones personales no demostrables.", bold_lead="Aportes rechazados o no incorporados.")
    add_para(doc, "Verificación. Revisión dimensional y analítica, pruebas unitarias, verificación de tipos, análisis estático, compilación productiva y contraste de PVGIS con PVWatts bajo parámetros comparables.", bold_lead="Verificación.")
    add_callout(doc, "Responsabilidad académica", "El grupo debe comprender, revisar y poder explicar todas las fórmulas, decisiones, resultados y limitaciones. La IA no sustituye la sustentación ni la autoría estudiantil.", fill=PALE_GOLD, accent=GOLD)
    add_editable(doc, "anexar o enlazar los prompts y respuestas relevantes que el profesor solicite; indicar quién revisó cada aporte y qué cambió.")

    doc.add_paragraph("21. Declaración de contribución de los integrantes", style="Heading 2")
    add_para(doc, "No se asignan contribuciones por inferencia. Cada integrante debe registrar únicamente tareas que pueda evidenciar y explicar.")
    add_table(
        doc,
        ["Integrante y código", "Contribución verificable", "Evidencia", "Revisión cruzada"],
        [
            ("[Nombre 1 - código]", "[Completar]", "[Archivo/acta/commit]", "[Quién revisó]"),
            ("[Nombre 2 - código]", "[Completar]", "[Archivo/acta/commit]", "[Quién revisó]"),
            ("[Nombre 3 - código]", "[Completar]", "[Archivo/acta/commit]", "[Quién revisó]"),
        ],
        [2200, 3200, 2200, 1760],
        font_size=8.8,
    )
    add_editable(doc, "agregar filas si existen más integrantes y escribir una afirmación breve de corresponsabilidad: todos revisaron el informe, conocen el proyecto completo y están en capacidad de sustentarlo.")

    add_section_title(doc, "Plan de comunicación del Hito 3", page_break=False)
    doc.add_paragraph("22. Póster académico", style="Heading 2")
    add_para(doc, "La tercera entrega exige un póster académico virtual o impreso. Debe permitir comprender el problema, el método, los resultados de los casos calibrados y su contraste sin leer el informe completo. La información debe coincidir con la versión final de este documento.")
    for text in (
        "Título y pregunta de investigación.",
        "Contexto y alcance: conectado a red, sin baterías, prefactibilidad.",
        "Entradas del caso y fuentes trazables.",
        "Dos ecuaciones centrales y explicación del redondeo.",
        "Resultado principal del caso y restricción de techo.",
        "Gráfica mensual, función n(Y) y significado de n'(Y).",
        "Contraste PVGIS/PVWatts y límites de la simulación.",
        "Limitaciones, declaración de IA y autores.",
    ):
        add_step(doc, text, poster_id)

    doc.add_paragraph("23. Pitch de 3 minutos", style="Heading 2")
    add_callout(doc, "Regla oficial", "El pitch se apoya exclusivamente en el póster. No se permiten PowerPoint, Canva ni otras diapositivas durante la presentación.", fill=PALE_GOLD, accent=GOLD)
    add_table(
        doc,
        ["Tiempo", "Contenido", "Pregunta que responde"],
        [
            ("0:00-0:35", "Problema, usuario y pregunta.", "¿Qué problema abordamos?"),
            ("0:35-1:35", "Datos, ecuaciones, derivada y algoritmo.", "¿Cómo lo resolvimos?"),
            ("1:35-2:20", "Resultado del caso, gráfica y techo.", "¿Qué encontramos?"),
            ("2:20-2:50", "Contraste, supuestos y limitaciones.", "¿Qué respalda y limita el resultado?"),
            ("2:50-3:00", "Cierre y propuesta de valor.", "¿Por qué importa?"),
        ],
        [1600, 4400, 3360],
        font_size=9.0,
    )
    add_editable(doc, "asignar intervenciones reales, ensayar con cronómetro y registrar una versión de máximo 3:00 minutos.")

    add_section_title(doc, "Referencias", page_break=True)
    add_para(doc, "Las referencias se presentan como base verificable. El grupo debe conservar los archivos consultados y completar la ficha del panel y los soportes específicos del caso.")
    references = (
        "Área Académica de Ciencias Básicas, Biológicas y Ambientales. (2026). Rúbrica de evaluación - Proyectos PjBL en Matemáticas [Documento del curso].",
        "Coordinación de Cálculo Diferencial, Cálculo Vectorial y Álgebra Lineal. (2026). Proyecto CAP+IA: instructivo [Documento del curso].",
        "Coordinación de Cálculo Diferencial, Cálculo Vectorial y Álgebra Lineal. (2026). Rúbrica de evaluación - tercera y última entrega de proyecto [Documento del curso].",
        "Universidad de Bogotá Jorge Tadeo Lozano. (2026). Sílabo de la asignatura Cálculo Diferencial, grupo 12, periodo 2026-2S [Documento del curso].",
        "Barragán, S. P., Melo, J., y Aya, O. (2021). Álgebra lineal: modelación, solución de problemas y ejercicios. Universidad de Bogotá Jorge Tadeo Lozano. https://doi.org/10.21789/9789587253177",
        "Stewart, J. (2012). Cálculo de una variable: trascendentes tempranas (7.ª ed.). Cengage Learning.",
        "European Commission, Joint Research Centre. (s. f.). Photovoltaic Geographical Information System (PVGIS). https://joint-research-centre.ec.europa.eu/photovoltaic-geographical-information-system-pvgis_en",
        "European Commission, Joint Research Centre. (s. f.). PVGIS non-interactive service. https://joint-research-centre.ec.europa.eu/photovoltaic-geographical-information-system-pvgis/using-pvgis-5/api-non-interactive-service_en",
        "European Commission, Joint Research Centre. (s. f.). PVGIS 5 user manual. https://joint-research-centre.ec.europa.eu/photovoltaic-geographical-information-system-pvgis/using-pvgis-5/pvgis-5-user-manual_en",
        "National Laboratory of the Rockies. (s. f.). PVWatts V8 API. https://developer.nlr.gov/docs/solar/pvwatts/v8/",
        "Unidad de Planeación Minero Energética y CORPOEMA. (2019). Estimación de los consumos de subsistencia en energía eléctrica, gas natural y GLP en territorio nacional SIN y ZNI: informe final. https://bdigital.upme.gov.co/handle/001/1367",
        "LONGi Solar. (s. f.). Hi-MO X6 Explorer LR5-54HTH 445-455M: ficha técnica oficial del módulo fotovoltaico. https://static.longi.com/Hi_MO_X6_LR_5_54_HTH_445_455_M_V1_30_30_and_15_V19_ES_09bf28ff6d.pdf",
        "Equipo SolarFit Hogar. (2026). Código fuente y documentación técnica del prototipo SolarFit Hogar [Repositorio local del proyecto].",
    )
    for ref in references:
        add_step(doc, ref, references_id)
    add_editable(doc, "agregar cualquier normativa o fuente adicional que el grupo use efectivamente y conservar copia de las respuestas de los proveedores solares.")

    # ANNEXES
    add_section_title(doc, "Anexo A - Ficha para un caso propio opcional", page_break=True)
    add_para(doc, "Los cinco casos calibrados ya documentan la prueba del MVP. Esta ficha permite que el grupo o un compañero agregue después una vivienda propia; sus datos no deben mezclarse con los casos oficiales calibrados.")
    add_table(
        doc,
        ["Campo", "Registro del grupo", "Soporte"],
        [
            ("Caso anonimizado", "[Municipio y descripción general]", "[Acta/criterio de selección]"),
            ("Coordenadas generales", "[Latitud, longitud]", "[Método de obtención]"),
            ("Área útil", "[___] m²", "[Medición, croquis o fotografía]"),
            ("Panel", "[Fabricante y modelo]", "[Ficha técnica oficial]"),
            ("Meta q", "[___] %", "[Justificación]"),
            ("Fuente solar", "[PVGIS/otra]", "[Respuesta y fecha]"),
            ("Fuente de contraste", "[PVWatts/otra]", "[Respuesta y fecha]"),
        ],
        [2700, 3360, 3300],
        font_size=9.0,
    )
    add_para(doc, "Registro mensual:", bold_lead="Registro mensual:")
    month_names = ("Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic")
    add_table(
        doc,
        ["Mes", "Consumo kWh", "Días", "HSP o Yₘ neto", "Fuente/observación"],
        [(month, "[___]", "[___]", "[___]", "[___]") for month in month_names],
        [900, 1760, 1200, 2200, 3300],
        font_size=8.5,
    )

    add_section_title(doc, "Anexo B - Protocolo reproducible", page_break=True)
    for text in (
        "Elegir ruta: facturas anonimizadas o simulación calibrada con una estadística oficial; documentar la decisión.",
        "Si un periodo cruza meses, documentar el prorrateo diario aplicado.",
        "Medir el área útil y descontar obstáculos, bordes y zonas no utilizables.",
        "Elegir un panel real y guardar la ficha técnica oficial.",
        "Definir meta y estrategia con una justificación explícita.",
        "Consultar PVGIS con potencia normalizada de 1 kWp y guardar configuración y respuesta.",
        "Ejecutar el cálculo; revisar solución continua, redondeo, cobertura y advertencias.",
        "Consultar una fuente independiente con parámetros comparables.",
        "Contrastar rendimiento y decisión de paneles; explicar las discrepancias sin atribuir exactitud empírica.",
        "Actualizar informe, póster y pitch con exactamente los mismos resultados.",
        "Hacer revisión cruzada: cada integrante explica una sección que no lideró.",
    ):
        add_step(doc, text, protocol_id)
    add_callout(doc, "Regla de trazabilidad", "Cada número del resultado final debe poder vincularse con una fuente, una transformación y una unidad. Si no puede reconstruirse, no debe presentarse como evidencia.")

    add_section_title(doc, "Anexo C - Matriz de verificación académica", page_break=True)
    add_table(
        doc,
        ["Requisito", "Ubicación en el informe", "Estado antes de entrega"],
        [
            ("Contextualización y objetivos", "Secciones 1-2", "Desarrollado"),
            ("Información previa y fuentes", "Sección 4 y Referencias", "Desarrollado para cinco casos"),
            ("Primer modelo", "Sección 5", "Desarrollado"),
            ("Desarrollo matemático", "Sección 7", "Desarrollado; grupo debe sustentarlo"),
            ("Resultados intermedios", "Secciones 11-12", "Simulación calibrada rotulada"),
            ("Visualización/simulación", "Figura 1 y prototipo", "Implementado"),
            ("Ajustes y revisión", "Sección 11", "Desarrollado"),
            ("Producto funcional", "Sección 13", "Verificado internamente"),
            ("Contraste de casos", "Secciones 14-15", "PVGIS/PVWatts ejecutado"),
            ("Reflexión crítica", "Sección 18", "Borrador para revisión personal"),
            ("IA y contribuciones", "Secciones 20-21", "Contribuciones pendientes"),
            ("Póster y pitch", "Secciones 22-23", "Preparar con datos finales"),
        ],
        [3200, 3000, 3160],
        font_size=8.8,
    )
    add_callout(doc, "Condición de cierre", "El contenido técnico y los cinco casos están desarrollados. Antes de entregar, el grupo debe completar nombres y aportes reales, comprobar que el póster muestre las mismas cifras y asegurarse de que todos puedan explicar el proyecto.", fill=PALE_GOLD, accent=GOLD)

    add_section_title(doc, "Anexo D - Evidencia técnica de software", page_break=True)
    add_para(doc, "Comando ejecutado el 7 de septiembre de 2026:")
    add_equation(doc, "pnpm check")
    add_table(
        doc,
        ["Etapa", "Resultado observado", "Alcance"],
        [
            ("ESLint", "Sin errores", "Calidad estática del código"),
            ("TypeScript", "Sin errores de tipos", "Coherencia de contratos"),
            ("Vitest", "9 archivos; 38 pruebas aprobadas", "Motor, archivos, persistencia, proveedores, interfaz, contraste e inclinaciones"),
            ("Next build", "Compilación productiva exitosa", "Página estática y tres familias de rutas API"),
        ],
        [2100, 3000, 4260],
        font_size=9.0,
    )
    add_para(doc, "Las pruebas cubren redondeo al final, no duplicación de pérdidas, techo inviable, monotonicidad de n(Y), criterio conservador, primera derivada, prorrateo de facturas, normalización de rendimientos, contraste a igual capacidad y comparación de inclinaciones.")
    add_callout(doc, "Alcance de esta evidencia", "Una compilación correcta y 38 pruebas aprobadas demuestran consistencia interna del prototipo en los casos programados. No verifican por sí solas disponibilidad futura de APIs externas, producción física, sombras del sitio ni viabilidad normativa.")

    add_section_title(doc, "Anexo E - Lista de cierre del grupo", page_break=True)
    checklist = (
        "[ ] Nombres y códigos completos en portada y declaración de contribuciones.",
        "[ ] Los cinco casos se presentan como simulaciones calibradas, no como viviendas medidas.",
        "[ ] Periodos y días revisados; prorrateo explicado cuando corresponde.",
        "[ ] Área útil, método de medición y obstáculos documentados.",
        "[ ] Ficha oficial del panel archivada y citada.",
        "[ ] PVGIS ejecutado con configuración final; respuesta conservada.",
        "[ ] Fuente independiente ejecutada con supuestos comparables.",
        "[ ] PVGIS y PVWatts se contrastan usando la misma capacidad instalable y parámetros comparables.",
        "[ ] Resultados de los cinco casos coinciden con los archivos de evidencia.",
        "[ ] Promedios oficiales, estimaciones solares y supuestos aparecen diferenciados.",
        "[ ] Póster y documento muestran cifras idénticas.",
        "[ ] Pitch cronometrado en 3 minutos y apoyado solo en el póster.",
        "[ ] Cada integrante puede explicar n(Y), n'(Y), algoritmo, contraste y límites.",
        "[ ] Declaración de IA revisada; prompts/evidencias anexados si se solicitan.",
        "[ ] Fecha operativa confirmada en AVATA.",
    )
    for item in checklist:
        add_para(doc, item, align=WD_ALIGN_PARAGRAPH.LEFT, after=5)

    add_callout(doc, "Cierre", "SolarFit Hogar está construido como producto académico y base de una posible startup. La calidad final de la entrega depende ahora de que el grupo cierre la evidencia del caso, revise cada afirmación y haga propia la sustentación.")

    # Keep headings with the next paragraph and prevent orphaned final lines where possible.
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.widow_control = True

    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
