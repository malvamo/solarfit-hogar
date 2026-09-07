from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_final_report as sf


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = PROJECT_ROOT.parents[2]
OUTPUT_DIR = WORKSPACE / "Entregables" / "CalculoDiferencial"
ASSET_DIR = PROJECT_ROOT / "docs" / "assets" / "semester-deliverables"
HITO2_OUTPUT = OUTPUT_DIR / "Hito_2_SolarFit_Hogar_Semana_10.docx"
FINAL_OUTPUT = OUTPUT_DIR / "Proyecto_Final_SolarFit_Hogar_Hito_3.docx"


CASES = [
    {
        "id": "SIM-BOG-001",
        "city": "Bogotá",
        "context": "Escenario residencial urbano de clima frío (>1.500 m)",
        "coords": "4,705583; −74,150667",
        "roof": 20,
        "consumption": [156] * 12,
        "pvgis": 1260.65,
        "pvwatts": 1271.0638046,
        "n_cont": 3.299885,
        "panels": 4,
        "pvwatts_panels": 4,
        "conservative_panels": 5,
        "buildable": 4,
        "max_roof": 9,
        "kwp": 1.80,
        "production": 2269.17,
        "coverage": 121.216,
        "min_coverage": 98.17,
        "mape": 4.653,
        "annual_diff": -0.819,
        "critical": "junio",
        "rule": "156 kWh/mes: promedio SUI 2018 reportado por UPME/CORPOEMA para usuarios del SIN ubicados por encima de 1.500 m s. n. m. [5, p. 28].",
    },
    {
        "id": "SIM-BAQ-002",
        "city": "Barranquilla",
        "context": "Escenario residencial urbano de clima cálido (<500 m)",
        "coords": "11,040000; −74,820833",
        "roof": 35,
        "consumption": [279] * 12,
        "pvgis": 1567.18,
        "pvwatts": 1643.3257448,
        "n_cont": 4.747381,
        "panels": 5,
        "pvwatts_panels": 5,
        "conservative_panels": 6,
        "buildable": 5,
        "max_roof": 16,
        "kwp": 2.25,
        "production": 3526.16,
        "coverage": 105.321,
        "min_coverage": 92.31,
        "mape": 4.741,
        "annual_diff": -4.634,
        "critical": "noviembre",
        "rule": "279 kWh/mes: promedio SUI 2018 reportado por UPME/CORPOEMA para usuarios del SIN ubicados por debajo de 500 m s. n. m. [5, p. 28].",
    },
    {
        "id": "SIM-MDE-003",
        "city": "Medellín",
        "context": "Escenario residencial urbano de clima templado (500–1.500 m)",
        "coords": "6,220611; −75,588972",
        "roof": 14,
        "consumption": [178] * 12,
        "pvgis": 1387.56,
        "pvwatts": 1369.6765144,
        "n_cont": 3.420873,
        "panels": 4,
        "pvwatts_panels": 4,
        "conservative_panels": 4,
        "buildable": 4,
        "max_roof": 6,
        "kwp": 1.80,
        "production": 2497.61,
        "coverage": 116.929,
        "min_coverage": 100.12,
        "mape": 4.427,
        "annual_diff": 1.306,
        "critical": "noviembre",
        "rule": "178 kWh/mes: promedio SUI 2018 reportado por UPME/CORPOEMA para usuarios del SIN ubicados entre 500 y 1.500 m s. n. m. [5, p. 28].",
    },
    {
        "id": "SIM-CLO-004",
        "city": "Cali",
        "context": "Escenario residencial urbano de clima templado (500–1.500 m)",
        "coords": "3,454500; −76,499722",
        "roof": 18,
        "consumption": [178] * 12,
        "pvgis": 1235.28,
        "pvwatts": 1422.1844213,
        "n_cont": 3.842584,
        "panels": 4,
        "pvwatts_panels": 4,
        "conservative_panels": 5,
        "buildable": 4,
        "max_roof": 8,
        "kwp": 1.80,
        "production": 2223.50,
        "coverage": 104.097,
        "min_coverage": 90.90,
        "mape": 13.167,
        "annual_diff": -13.142,
        "critical": "junio",
        "rule": "178 kWh/mes: promedio SUI 2018 reportado por UPME/CORPOEMA para usuarios del SIN ubicados entre 500 y 1.500 m s. n. m. [5, p. 28].",
    },
    {
        "id": "SIM-UIB-005",
        "city": "Quibdó",
        "context": "Escenario residencial rural de clima cálido con techo restringido",
        "coords": "5,690556; −76,643778",
        "roof": 8,
        "consumption": [118.8] * 12,
        "pvgis": 898.28,
        "pvwatts": 1131.8748349,
        "n_cont": 3.526740,
        "panels": 4,
        "pvwatts_panels": 3,
        "conservative_panels": 5,
        "buildable": 3,
        "max_roof": 3,
        "kwp": 1.80,
        "production": 1212.68,
        "coverage": 85.064,
        "min_coverage": 74.05,
        "mape": 20.386,
        "annual_diff": -20.638,
        "critical": "junio",
        "rule": "118,8 kWh/mes: promedio rural SUI 2018 reportado por UPME/CORPOEMA para clima cálido [5, p. 81]. El techo de 8 m² es un supuesto de prueba, no una medición.",
    },
]


REFERENCES = [
    ("[1]", "Rúbrica general de proyectos CAP+IA de Cálculo Diferencial, momentos 1–3, archivos oficiales del curso."),
    ("[2]", "Rúbrica de la tercera entrega, semana 14: documento acumulativo, póster y pitch de 3 minutos, archivos oficiales del curso."),
    ("[3]", "JRC, ‘PVGIS 5.3 API non-interactive service’. https://joint-research-centre.ec.europa.eu/photovoltaic-geographical-information-system-pvgis/using-pvgis-5/api-non-interactive-service_en"),
    ("[4]", "NLR, ‘PVWatts V8 API’. https://developer.nlr.gov/docs/solar/pvwatts/v8/"),
    ("[5]", "UPME y CORPOEMA (2019), ‘Estimación de los consumos de subsistencia en energía eléctrica, gas natural y GLP en territorio nacional SIN y ZNI: informe final’. Promedios SUI 2018 por piso térmico (p. 28) y sector rural (p. 81). https://bdigital.upme.gov.co/handle/001/1367"),
    ("[6]", "LONGi, ficha oficial LR5-54HTH-445–455M. https://static.longi.com/Hi_MO_X6_LR_5_54_HTH_445_455_M_V1_30_30_and_15_V19_ES_09bf28ff6d.pdf"),
    ("[7]", "CREG, Resolución 174 de 2021, versión consolidada. https://gestornormativo.creg.gov.co/gestor/entorno/docs/resolucion_creg_0174_2021.htm"),
    ("[8]", "MinEnergía, portal oficial RETIE. https://www.minenergia.gov.co/es/misional/energia-electrica-2/reglamentos-tecnicos/reglamento-t%C3%A9cnico-de-instalaciones-el%C3%A9ctricas-retie/"),
]


def annual_consumption(case: dict) -> float:
    return float(sum(case["consumption"]))


def format_kwh(value: float) -> str:
    decimals = 0 if value.is_integer() else 1
    return f"{value:,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _font(size: int, *, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def build_yield_chart(path: Path) -> None:
    width, height = 1800, 980
    img = Image.new("RGB", (width, height), "#F7F5EF")
    draw = ImageDraw.Draw(img)
    title = _font(58, bold=True)
    label = _font(34)
    small = _font(26)
    bold_small = _font(28, bold=True)
    draw.text((90, 55), "PVGIS y PVWatts no siempre coinciden", fill="#123F42", font=title)
    draw.text((90, 130), "Rendimiento específico anual · kWh/kWp·año · misma ubicación y supuestos", fill="#52666A", font=small)
    left, top, right, bottom = 150, 245, 1700, 805
    max_value = 1800
    for tick in range(0, max_value + 1, 300):
        y = bottom - (tick / max_value) * (bottom - top)
        draw.line((left, y, right, y), fill="#CBD7D4", width=2)
        draw.text((45, y - 17), f"{tick}", fill="#52666A", font=small)
    group_w = (right - left) / len(CASES)
    bar_w = 74
    for index, case in enumerate(CASES):
        center = left + group_w * (index + 0.5)
        values = [(case["pvgis"], "#1E7B78"), (case["pvwatts"], "#F4B63D")]
        for j, (value, color) in enumerate(values):
            x1 = center + (-bar_w - 8 if j == 0 else 8)
            y1 = bottom - value / max_value * (bottom - top)
            draw.rounded_rectangle((x1, y1, x1 + bar_w, bottom), radius=12, fill=color)
        draw.text((center - 75, bottom + 28), case["city"], fill="#123F42", font=bold_small)
        draw.text((center - 64, bottom + 70), f"{case['annual_diff']:+.1f}%", fill="#B3483D" if abs(case["annual_diff"]) >= 10 else "#52666A", font=small)
    draw.rounded_rectangle((1250, 65, 1300, 110), radius=8, fill="#1E7B78")
    draw.text((1315, 70), "PVGIS 5.3", fill="#123F42", font=small)
    draw.rounded_rectangle((1510, 65, 1560, 110), radius=8, fill="#F4B63D")
    draw.text((1575, 70), "PVWatts V8", fill="#123F42", font=small)
    draw.text((90, 905), "La cifra inferior es la diferencia anual PVGIS respecto de PVWatts. Cali y Quibdó exigen cautela.", fill="#52666A", font=label)
    img.save(path, quality=95)


def build_coverage_chart(path: Path) -> None:
    width, height = 1800, 980
    img = Image.new("RGB", (width, height), "#F7F5EF")
    draw = ImageDraw.Draw(img)
    title = _font(58, bold=True)
    small = _font(28)
    bold_small = _font(30, bold=True)
    draw.text((90, 55), "Cuatro casos cumplen; Quibdó revela la restricción", fill="#123F42", font=title)
    draw.text((90, 132), "Cobertura anual calculada con la cantidad construible por área", fill="#52666A", font=small)
    left, top, right, bottom = 170, 245, 1700, 800
    max_value = 140
    target_y = bottom - 100 / max_value * (bottom - top)
    draw.line((left, target_y, right, target_y), fill="#B3483D", width=5)
    draw.text((left + 8, target_y - 42), "meta 100 %", fill="#B3483D", font=small)
    group_w = (right - left) / len(CASES)
    bar_w = 150
    for index, case in enumerate(CASES):
        center = left + group_w * (index + 0.5)
        value = case["coverage"]
        y1 = bottom - value / max_value * (bottom - top)
        color = "#1E7B78" if value >= 100 else "#D86D59"
        draw.rounded_rectangle((center - bar_w / 2, y1, center + bar_w / 2, bottom), radius=18, fill=color)
        draw.text((center - 58, y1 - 45), f"{value:.1f}%", fill="#123F42", font=bold_small)
        draw.text((center - 75, bottom + 28), case["city"], fill="#123F42", font=bold_small)
        draw.text((center - 65, bottom + 72), f"{case['buildable']}/{case['panels']} paneles", fill="#52666A", font=small)
    draw.text((90, 905), "Paneles usados/requeridos. El modelo conserva 4 requeridos en Quibdó y advierte que solo caben 3.", fill="#52666A", font=small)
    img.save(path, quality=95)


def add_figure(doc: Document, image_path: Path, caption: str, alt_text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Inches(6.25))
    shape._inline.docPr.set("descr", alt_text)
    c = doc.add_paragraph(caption, style="Caption")
    c.paragraph_format.keep_with_next = False


def new_document(title: str, subtitle: str, label: str) -> Document:
    doc = Document()
    sf.configure_styles(doc)
    doc.settings.odd_and_even_pages_header_footer = False
    sf.configure_cover_section(doc.sections[0])
    doc.core_properties.title = title
    doc.core_properties.subject = "Proyecto integrador CAP+IA de Cálculo Diferencial"
    doc.core_properties.author = "Equipo SolarFit Hogar"
    doc.core_properties.keywords = "cálculo diferencial, energía solar, simulación, PVGIS, PVWatts, CAP+IA"
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(11)
    p = doc.add_paragraph("PROYECTO INTEGRADOR CAP+IA", style="SolarFit Kicker")
    p.paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph("SolarFit Hogar", style="SolarFit Cover Title")
    p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph(subtitle, style="SolarFit Cover Subtitle")
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph(label, style="SolarFit Kicker")
    p.paragraph_format.space_after = Pt(42)
    for text, size, color, bold in (
        ("Cálculo Diferencial · Código 502116 · Grupo 12", 11.5, sf.TEAL, True),
        ("Fundación Universidad de Bogotá Jorge Tadeo Lozano", 11, sf.INK, False),
        ("Docente: Orlando Aya Corredor", 10.5, sf.MUTED, False),
        ("Periodo académico: 2026-2S", 10.5, sf.MUTED, False),
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        sf.set_run_font(paragraph.add_run(text), size=size, color=color, bold=bold)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(24)
    sf.set_run_font(paragraph.add_run("Integrantes y códigos: [COMPLETAR ANTES DE ENTREGAR]"), size=10.5, color=sf.RED, bold=True)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(14)
    sf.set_run_font(paragraph.add_run("Versión académica corregida: 7 de septiembre de 2026"), size=9.5, color=sf.MUTED, italic=True)
    content = doc.add_section(WD_SECTION.NEW_PAGE)
    content.header.is_linked_to_previous = False
    content.footer.is_linked_to_previous = False
    sf.configure_content_section(content)
    return doc


def add_case_summary_table(doc: Document, *, include_conservative: bool = False) -> None:
    headers = ["Caso", "Demanda anual", "PVGIS", "Paneles", "Cobertura / estado"]
    rows = []
    for case in CASES:
        panels = f"{case['panels']} anual"
        if include_conservative:
            panels += f" · {case['conservative_panels']} conserv."
        state = f"{case['coverage']:.2f} %"
        if case["buildable"] < case["panels"]:
            state += f" · REVISAR ({case['buildable']} caben)"
        rows.append((
            f"{case['id']}\n{case['city']}",
            f"{format_kwh(annual_consumption(case))} kWh",
            f"{case['pvgis']:,.2f} kWh/kWp·año".replace(",", "X").replace(".", ",").replace("X", "."),
            panels,
            state.replace(".", ","),
        ))
    sf.add_table(doc, headers, rows, [1900, 1500, 2100, 1700, 2160], font_size=8.7)


def add_references(doc: Document) -> None:
    for key, value in REFERENCES:
        sf.add_para(doc, f"{key} {value}", align=WD_ALIGN_PARAGRAPH.LEFT, after=4)


def add_identity_and_ai(doc: Document) -> None:
    sf.add_section_title(doc, "Declaración de IA y contribuciones", page_break=True)
    sf.add_callout(
        doc,
        "Uso de IA",
        "ChatGPT/Codex se utilizó para analizar requisitos, investigar fuentes oficiales, estructurar el modelo, programar y probar el MVP, ejecutar los escenarios, calcular comparaciones, redactar y diseñar borradores. El grupo conserva la responsabilidad de verificar cifras, revisar fuentes, comprender el modelo y defender oralmente el trabajo.",
        fill=sf.PALE_GOLD,
        accent=sf.GOLD,
    )
    sf.add_editable(doc, "escribir nombres, códigos y aporte real verificable de cada integrante; no asignar aportes ficticios ni idénticos si no ocurrieron.")
    sf.add_table(
        doc,
        ["Integrante", "Aporte real", "Evidencia / revisión"],
        [
            ("[Nombre 1]", "[Completar]", "[Archivo, prueba o sección revisada]"),
            ("[Nombre 2]", "[Completar]", "[Archivo, prueba o sección revisada]"),
            ("[Nombre 3]", "[Completar]", "[Archivo, prueba o sección revisada]"),
            ("[Nombre 4, si aplica]", "[Completar]", "[Archivo, prueba o sección revisada]"),
        ],
        [1900, 3800, 3660],
        font_size=9.0,
    )


def finalize_doc(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.widow_control = True


def build_hito2() -> Document:
    doc = new_document(
        "SolarFit Hogar - Hito 2",
        "Desarrollo matemático, simulación multicase y revisión del modelo",
        "HITO 2 · INVESTIGACIÓN Y CREACIÓN · SEMANA 10",
    )
    sf.add_section_title(doc, "Ficha de entrega")
    sf.add_callout(
        doc,
        "Estado",
        "Informe de avance cerrado con cinco casos simulados calibrados con fuentes oficiales. Las series de consumo no son facturas ni observaciones de hogares; PVGIS y PVWatts entregan simulaciones solares, no producción medida.",
        fill=sf.PALE_GOLD,
        accent=sf.GOLD,
    )
    sf.add_table(
        doc,
        ["Elemento", "Evidencia incluida", "Estado"],
        [
            ("Desarrollo matemático", "Balance energético y sensibilidad directa n(Y), n′(Y)", "Cerrado"),
            ("Resultados intermedios", "Cinco ubicaciones y dos criterios de dimensionamiento", "Cerrado"),
            ("Visualización / simulación", "Dos gráficas y estimaciones PVGIS/PVWatts", "Cerrado"),
            ("Ajustes del modelo", "Pérdidas sin doble conteo, techo no truncado, comparación externa", "Cerrado"),
            ("Identidad y aportes", "Debe completarlos el grupo", "Pendiente humano"),
        ],
        [2300, 4760, 2300],
        font_size=9.0,
    )

    sf.add_section_title(doc, "1. Evolución desde el Hito 1")
    sf.add_para(doc, "El Hito 1 formuló la pregunta central: cuántos paneles necesita una vivienda para compensar su consumo. El Hito 2 transforma esa pregunta en un modelo ejecutable, distingue potencia de energía, incorpora doce periodos, usa rendimiento solar mensual y separa la recomendación energética de la restricción física del techo.")
    sf.add_table(
        doc,
        ["Versión inicial", "Ajuste realizado", "Razón matemática o técnica"],
        [
            ("Consumo ÷ potencia", "Consumo objetivo ÷ producción específica", "kWh no se divide directamente por kW; se necesita energía por kWp."),
            ("Un promedio solar", "Doce rendimientos mensuales", "Permite reconocer variación y mes limitante."),
            ("PR aplicado siempre", "Modos HSP y PVGIS excluyentes", "PVGIS ya entrega energía con pérdidas configuradas."),
            ("Paneles limitados por techo", "Recomendado y construible separados", "Una restricción no debe ocultar la necesidad matemática."),
            ("Una herramienta", "PVGIS principal + PVWatts contraste", "Dos modelos revelan sensibilidad a bases y algoritmos distintos."),
        ],
        [2400, 2860, 4100],
        font_size=9.0,
    )

    sf.add_section_title(doc, "2. Datos y protocolo de simulación", page_break=True)
    sf.add_para(doc, "Los cinco escenarios son simulaciones calibradas, no viviendas observadas. La demanda mensual se fija con promedios SUI 2018 publicados por UPME/CORPOEMA para el piso térmico o sector rural correspondiente, sin inventar estacionalidad que la fuente no documenta. El recurso solar se estima con PVGIS y se contrasta con PVWatts. Las coordenadas y el área del techo son supuestos de prueba identificados expresamente.")
    sf.add_table(
        doc,
        ["Componente", "Dato confirmado", "Supuesto / transformación", "Impacto"],
        [
            ("Panel", "LONGi LR5-54HTH-450M: 450 Wp; 1,722 × 1,134 m [6]", "Factor de instalación 1,10", "Área reservada: 2,148 m² por panel."),
            ("Demanda", "Promedios SUI 2018 por piso térmico y sector rural [5]", "El promedio se repite 12 meses", "Calibra el caso, pero no representa una factura ni un promedio municipal."),
            ("Ubicación", "Sin vivienda o estación observada", "Punto geográfico de referencia", "No modela sombra ni geometría local."),
            ("Producción", "PVGIS 5.3 y PVWatts V8 [3][4]", "10°, sur, 14 % pérdidas", "Contraste intermodelo, no medición."),
            ("Techo", "Ninguna medición física", "8–35 m² útiles según caso", "Prueba lógica de viabilidad preliminar."),
        ],
        [1500, 2940, 2450, 2470],
        font_size=8.7,
    )
    sf.add_para(doc, "Consumos enero–diciembre:", bold_lead="Consumos enero–diciembre:", keep_with_next=True)
    for case in CASES:
        sf.add_para(doc, f"{case['id']} {case['city']}: " + ", ".join(str(v) for v in case["consumption"]) + " kWh. " + case["rule"], align=WD_ALIGN_PARAGRAPH.LEFT, after=4)

    sf.add_section_title(doc, "3. Desarrollo matemático", page_break=True)
    sf.add_para(doc, "Para PVGIS, Yₘ es el rendimiento específico neto del mes m en kWh/kWp. Si P es la potencia de un panel en kWp, q la meta y Cₘ el consumo, entonces:")
    sf.add_equation(doc, "Gₘ = P·Yₘ       y       nₘ* = q·Cₘ / Gₘ")
    sf.add_para(doc, "El criterio anual compensa energía total, mientras el criterio conservador pretende alcanzar la meta en cada mes:")
    sf.add_equation(doc, "n*anual = q·ΣCₘ / (P·ΣYₘ)       n*conservador = máxₘ(q·Cₘ / (P·Yₘ))")
    sf.add_equation(doc, "N = ceil(n*)       Nmáx = floor(Aᵤ /(ancho·alto·fA))")
    sf.add_para(doc, "N se redondea una sola vez al final. Si N > Nmáx, el sistema mantiene N como necesidad energética y calcula por separado Nconstruible = Nmáx.")
    doc.add_paragraph("3.1 Sensibilidad que sí afecta la decisión", style="Heading 2")
    sf.add_para(doc, "Sea D=ΣCₘ la demanda anual y Y=ΣYₘ el rendimiento específico anual. Antes de redondear, la necesidad de paneles como función del recurso solar es:")
    sf.add_equation(doc, "n(Y)=q·D/(P·Y)       n′(Y)=−q·D/(P·Y²)")
    sf.add_para(doc, "Como n′(Y)<0, un mayor rendimiento solar reduce la necesidad continua de paneles. El valor de n′(Y) cuantifica el cambio local esperado ante una variación pequeña de Y y ayuda a explicar por qué la recomendación puede ser más sensible en una ubicación que en otra. Esta relación tiene impacto directo en el dimensionamiento; no se agregan procedimientos matemáticos que no cambien la decisión del proyecto.")
    sf.add_callout(doc, "Aplicación directa", "En Bogotá, n′(1.260,65)=−0,00262 paneles por cada kWh/kWp·año. Cerca de ese escenario, una diferencia de 100 kWh/kWp·año cambia la necesidad continua en aproximadamente 0,262 paneles, antes del redondeo. Por eso contrastar el recurso solar puede afectar la recomendación entera.")
    doc.add_paragraph("3.2 Comparación discreta de inclinaciones", style="Heading 2")
    sf.add_para(doc, "La aplicación consulta un conjunto finito de inclinaciones y compara el rendimiento anual devuelto por PVGIS. Reporta la alternativa de mayor rendimiento dentro de ese conjunto, sin afirmar que sea un máximo exacto ni usar esa comparación como un ejercicio adicional de derivación.")

    sf.add_section_title(doc, "4. Representación algorítmica")
    steps = [
        "Validar doce consumos, unidades y parámetros globales.",
        "Consultar PVGIS para 1 kWp y conservar la serie mensual neta y sus metadatos.",
        "Calcular Gₘ y nₘ*; sumar para el criterio anual o tomar el máximo para el conservador.",
        "Redondear n* hacia arriba una sola vez para obtener N.",
        "Calcular Nmáx con el área útil y separar recomendación de cantidad construible.",
        "Simular producción, cobertura y balance de cada mes.",
        "Consultar PVWatts con supuestos comparables y contrastar rendimiento anual y decisión entera.",
        "Exportar configuración, series, resultados y hashes para reproducibilidad.",
    ]
    for item in steps:
        sf.add_bullet(doc, item, 0)
    sf.add_callout(doc, "Implementación", "El mismo motor TypeScript utilizado por la aplicación web procesa los escenarios; no se usó una calculadora paralela para construir las conclusiones. La versión corregida aprobó 38/38 pruebas automatizadas, ESLint, TypeScript y build productivo el 7 de septiembre de 2026.")

    sf.add_section_title(doc, "5. Resultados intermedios de cinco casos", page_break=True)
    add_case_summary_table(doc, include_conservative=True)
    sf.add_table_source(doc, "Fuente: ejecución reproducible del MVP el 7 de septiembre de 2026. Demanda calibrada con promedios oficiales agregados; producción estimada. La cobertura de Quibdó usa los 3 paneles construibles, no los 4 requeridos.")
    add_figure(doc, ASSET_DIR / "coverage-cases.png", "Figura 1. Cobertura anual con la cantidad construible. Cuatro casos superan la meta anual; Quibdó queda limitado por área.", "Gráfico de barras de cobertura anual para Bogotá, Barranquilla, Medellín, Cali y Quibdó; Quibdó alcanza 85,06 por ciento por la restricción de techo.")
    sf.add_callout(doc, "Lectura principal", "El criterio anual puede dejar meses por debajo de 100 %. Barranquilla obtiene 105,32 % anual, pero su mínimo mensual es 92,31 % en noviembre; el criterio conservador aumenta de 5 a 6 paneles.")

    sf.add_section_title(doc, "6. Contraste PVGIS–PVWatts", page_break=True)
    sf.add_table(
        doc,
        ["Ciudad", "PVGIS", "PVWatts", "Diferencia anual", "Paneles P / W"],
        [(c["city"], f"{c['pvgis']:.2f}", f"{c['pvwatts']:.2f}", f"{c['annual_diff']:+.2f} %", f"{c['panels']} / {c['pvwatts_panels']}") for c in CASES],
        [1650, 1700, 1700, 2100, 2210],
        font_size=9.1,
    )
    add_figure(doc, ASSET_DIR / "yield-comparison.png", "Figura 2. Rendimiento específico anual calculado por PVGIS y PVWatts bajo una configuración comparable.", "Barras agrupadas que comparan el rendimiento anual de PVGIS y PVWatts en cinco ciudades; las mayores diferencias aparecen en Cali y Quibdó.")
    sf.add_para(doc, "Bogotá, Barranquilla, Medellín y Cali conservan la misma decisión entera con ambos proveedores. En Quibdó, PVGIS recomienda 4 paneles y PVWatts 3. No se elige automáticamente el valor más favorable: la divergencia se reporta como incertidumbre de modelo.")
    sf.add_callout(doc, "Interpretación correcta", "PVGIS y PVWatts no son verdad de campo. Usan bases meteorológicas y modelos diferentes. Su comparación comprueba orden de magnitud, detecta sensibilidad y evita una confianza falsa en una sola fuente.")

    sf.add_section_title(doc, "7. Importancia académica de las simulaciones")
    for text in (
        "Permiten probar de forma reproducible cinco combinaciones de clima, demanda y área sin fabricar facturas ni patrones mensuales.",
        "Ejercitan casos favorables y uno adverso; el caso de Quibdó muestra que el software advierte una restricción de área.",
        "Conectan n′(Y) con una decisión: su signo negativo y su magnitud explican cómo cambia la necesidad continua cuando varía el recurso solar.",
        "Permiten comparar criterios anual y conservador sin esperar doce meses de una instalación.",
        "Conservan entradas y proveedores para que otro equipo repita la ejecución.",
    ):
        sf.add_bullet(doc, text, 0)
    sf.add_para(doc, "Su límite es igualmente importante: una simulación no valida sombra local, estructura, conexión eléctrica, tarifa ni desempeño medido. La rúbrica permite condiciones simuladas, pero no autoriza llamarlas evidencia empírica.")

    sf.add_section_title(doc, "8. Conclusiones del Hito 2")
    conclusions = [
        "El modelo ya es dimensionalmente coherente y ejecutable con doce periodos.",
        "La derivada n′(Y) explica la dirección y sensibilidad de la necesidad continua frente al rendimiento solar.",
        "Cuatro escenarios satisfacen la meta anual con el techo supuesto; uno activa correctamente la advertencia de área.",
        "El contraste entre simuladores descubre divergencias materiales en Cali y Quibdó que deben conservarse como riesgo.",
        "El avance cumple la función del Hito 2: desarrollo matemático, simulación, resultados intermedios y ajustes justificables.",
    ]
    for item in conclusions:
        sf.add_bullet(doc, item, 0)
    add_identity_and_ai(doc)
    sf.add_section_title(doc, "Referencias", page_break=True)
    add_references(doc)
    finalize_doc(doc)
    return doc


def build_final() -> Document:
    doc = new_document(
        "SolarFit Hogar - Informe final acumulativo",
        "Prefactibilidad fotovoltaica residencial explicable y evaluada con cinco simulaciones calibradas",
        "INFORME FINAL ACUMULATIVO · HITOS 1, 2 Y 3 · SEMANA 14",
    )
    sf.add_section_title(doc, "Resumen ejecutivo")
    sf.add_callout(
        doc,
        "Resultado",
        "SolarFit Hogar es un MVP web desplegado que dimensiona paneles por balance anual o criterio conservador, verifica el límite preliminar del techo, explica sensibilidad mediante cálculo diferencial y contrasta PVGIS con PVWatts. Cinco escenarios simulados calibrados muestran cuatro casos que cumplen la meta anual y uno que activa una advertencia de área insuficiente.",
    )
    sf.add_para(doc, "El producto responde una pregunta de prefactibilidad: qué potencia y cuántos paneles podrían compensar una meta energética bajo entradas declaradas. No certifica una instalación, no garantiza factura cero y no sustituye el diseño profesional.")
    sf.add_table(
        doc,
        ["Entrega", "Contenido acumulado", "Evidencia"],
        [
            ("Hito 1", "Contexto, pregunta, objetivos, variables y modelo inicial", "Problema y alcance definidos"),
            ("Hito 2", "Sensibilidad directa, algoritmo, simulación y ajustes", "Cinco casos + PVGIS/PVWatts"),
            ("Hito 3", "Producto, contraste, reflexión y comunicación", "MVP en Vercel, póster y pitch"),
        ],
        [1700, 4000, 3660],
        font_size=9.2,
    )

    sf.add_section_title(doc, "1. Problema, objetivos y alcance", page_break=True)
    sf.add_para(doc, "Una vivienda no puede dimensionarse dividiendo su consumo por la potencia nominal de un panel. El consumo está en kWh; la potencia, en kWp; y la energía producida cambia con la ubicación, el mes, la inclinación, las pérdidas y el techo disponible. SolarFit integra esas variables y conserva trazabilidad.")
    doc.add_paragraph("1.1 Pregunta", style="Heading 2")
    sf.add_callout(doc, "Pregunta orientadora", "¿Cuántos paneles y qué potencia fotovoltaica nominal necesita una vivienda conectada a red para compensar una fracción de su consumo bajo doce periodos, recurso solar modelado y área útil declarada?")
    doc.add_paragraph("1.2 Objetivo general", style="Heading 2")
    sf.add_para(doc, "Desarrollar y evaluar un MVP de prefactibilidad que convierta consumos mensuales, recurso solar, características del módulo y área útil en una recomendación explicable de potencia, paneles, cobertura y viabilidad preliminar.")
    doc.add_paragraph("1.3 Objetivos específicos", style="Heading 2")
    for text in (
        "Formular un modelo matemático con variables, unidades, dominio, supuestos y dos criterios de dimensionamiento.",
        "Aplicar n(Y) y n′(Y) para explicar cómo el rendimiento solar cambia directamente la necesidad continua de paneles.",
        "Implementar el modelo en una aplicación web con validación, persistencia local y exportaciones.",
        "Evaluar cinco escenarios reproducibles y contrastar la fuente principal con un segundo modelo oficial.",
        "Comunicar resultados, incertidumbre, limitaciones y ruta de escalamiento sin confundir simulación con medición.",
    ):
        sf.add_bullet(doc, text, 0)
    doc.add_paragraph("1.4 Alcance", style="Heading 2")
    sf.add_para(doc, "El sistema estudia compensación energética de una vivienda conectada a red y sin baterías. No modela simultaneidad horaria, autoconsumo, tarifa, excedentes, estructura, sombras cercanas, cableado, inversor específico ni trámite de conexión.")

    sf.add_section_title(doc, "2. Fuentes, datos y supuestos", page_break=True)
    sf.add_table(
        doc,
        ["Entrada", "Fuente", "Tratamiento"],
        [
            ("Consumo", "UPME/CORPOEMA [5]", "Promedios agregados SUI 2018 repetidos 12 meses; nunca facturas reales."),
            ("Ubicación", "Supuesto de escenario", "Punto geográfico reproducible; no corresponde a una vivienda ni estación medida."),
            ("Producción principal", "JRC PVGIS 5.3 [3]", "Serie mensual neta para 1 kWp, inclinación 10°, sur, 14 % pérdidas."),
            ("Contraste", "NLR PVWatts V8 [4]", "Mismos parámetros comparables; base y modelos internos distintos."),
            ("Módulo", "LONGi LR5-54HTH-450M [6]", "450 Wp, 1,722 × 1,134 m; factor de área supuesto 1,10."),
            ("Techo", "Supuesto de escenario", "Área útil escalar; no prueba disposición geométrica ni estructura."),
        ],
        [1700, 2600, 5060],
        font_size=8.8,
    )
    sf.add_callout(doc, "Control de supuestos", "Dato confirmado: promedios agregados de consumo publicados, especificación del panel y respuestas de APIs. Supuesto: coordenada de referencia y techo de cada escenario residencial hipotético. Razón: no se dispone de doce facturas, medición del techo ni un sistema instalado. Impacto: los resultados verifican el comportamiento esperado del modelo en escenarios calibrados, no una vivienda real.", fill=sf.PALE_GOLD, accent=sf.GOLD)

    sf.add_section_title(doc, "3. Modelo matemático", page_break=True)
    sf.add_para(doc, "Con Cₘ en kWh, q como meta, P en kWp y Yₘ como rendimiento específico neto en kWh/kWp:")
    sf.add_equation(doc, "Gₘ=P·Yₘ      Tₘ=q·Cₘ      nₘ*=Tₘ/Gₘ")
    sf.add_equation(doc, "n*anual=q·ΣCₘ/(P·ΣYₘ)      n*conservador=máxₘ(q·Cₘ/(P·Yₘ))")
    sf.add_equation(doc, "N=ceil(n*)      Nmáx=floor(Aᵤ/(ancho·alto·fA))      Nconstruible=mín(N,Nmáx)")
    sf.add_para(doc, "Si se usa la entrada alternativa de horas solares, Yₘ se construye con hₘ·dₘ·PR. Los modos son excluyentes: cuando PVGIS ya entrega energía con pérdidas configuradas, PR no se aplica otra vez.")
    doc.add_paragraph("3.1 Sensibilidad diferencial con impacto directo", style="Heading 2")
    sf.add_para(doc, "Con D=ΣCₘ y Y=ΣYₘ, la necesidad continua anual se puede expresar únicamente en función del rendimiento específico:")
    sf.add_equation(doc, "n(Y)=q·D/(P·Y)       n′(Y)=−q·D/(P·Y²)")
    sf.add_para(doc, "n′(Y)<0 demuestra que, si la demanda, la meta y el panel permanecen constantes, un aumento del rendimiento solar reduce la necesidad continua. Su magnitud permite estimar y comparar sensibilidad antes del redondeo entero. Esta es la aplicación de cálculo diferencial conservada porque modifica la interpretación de la recomendación.")
    sf.add_callout(doc, "Interpretación numérica", "En Bogotá, n′(1.260,65)=−0,00262 paneles por cada kWh/kWp·año; cerca de ese punto, 100 kWh/kWp·año adicionales reducen la necesidad continua en aproximadamente 0,262 paneles. En Quibdó, la magnitud es mayor: |n′(898,28)|≈0,00393. La comparación ayuda a priorizar cautela donde el recurso estimado es menor.")
    doc.add_paragraph("3.2 Comparación discreta de inclinaciones", style="Heading 2")
    sf.add_para(doc, "El MVP consulta varias inclinaciones predeterminadas y ordena los rendimientos anuales estimados por PVGIS. Esta función es una comparación de escenarios: identifica la mejor alternativa evaluada, pero no afirma un máximo exacto ni añade operaciones diferenciales sin efecto necesario en el producto.")

    sf.add_section_title(doc, "4. Producto funcional y arquitectura", page_break=True)
    sf.add_para(doc, "El MVP está desplegado en https://solarfit-hogar.vercel.app y versionado en el repositorio privado malvamo/solarfit-hogar. El frontend recibe los datos; las rutas de servidor consultan proveedores externos sin exponer la clave de PVWatts; el dominio valida y calcula; el navegador conserva un borrador local.")
    sf.add_table(
        doc,
        ["Capa", "Responsabilidad", "Evidencia"],
        [
            ("Interfaz", "Doce consumos, ubicación, panel, techo, meta y estrategia", "Next.js 16 + React 19"),
            ("Dominio", "Unidades, criterios, redondeo, techo, balances y n′(Y)", "TypeScript puro con pruebas"),
            ("PVGIS", "Rendimiento mensual y comparación discreta de inclinaciones", "API JRC desde backend"),
            ("PVWatts", "Serie AC mensual independiente", "Clave NLR solo en Vercel"),
            ("Evidencia", "JSON/CSV/Markdown por escenario y libro de validación", "Entradas, series, métricas y hash"),
        ],
        [1600, 4020, 3740],
        font_size=9.0,
    )
    sf.add_callout(doc, "Pruebas", "La versión auditada aprobó ESLint, TypeScript, build productivo y 38 pruebas automatizadas. Estas pruebas verifican la implementación; no reemplazan la validación del contexto físico.")

    sf.add_section_title(doc, "5. Protocolo de los cinco casos")
    for case in CASES:
        sf.add_para(doc, f"{case['id']} — {case['city']}. {case['context']}. Coordenada de referencia {case['coords']}; techo útil supuesto {case['roof']} m²; {case['rule']}", align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    sf.add_para(doc, "Configuración común: panel LONGi de 450 Wp; factor de área 1,10; meta 100 %; criterio anual; 10° de inclinación; azimut brújula 180°; pérdidas 14 %; incertidumbre ±10 %; ejecución reproducible del 7 de septiembre de 2026.")

    sf.add_section_title(doc, "6. Resultados multicase")
    add_case_summary_table(doc, include_conservative=True)
    add_figure(doc, ASSET_DIR / "coverage-cases.png", "Figura 1. Cobertura anual construible en los cinco escenarios simulados.", "Gráfico de cobertura anual que muestra cuatro casos sobre 100 por ciento y Quibdó en 85,06 por ciento.")
    sf.add_para(doc, "El caso restringido de Quibdó es un resultado funcional, no un error. El modelo calcula 3,527 paneles equivalentes, recomienda 4, reconoce que el techo supuesto admite 3 y reporta 85,06 % de cobertura construible. No trunca silenciosamente la necesidad.")

    sf.add_section_title(doc, "7. Caso principal: Bogotá")
    bog = CASES[0]
    sf.add_para(doc, "Se selecciona Bogotá como caso principal del póster y el pitch porque utiliza el promedio SUI 2018 publicado para usuarios del SIN ubicados por encima de 1.500 m, facilita la reproducción manual y conserva un margen de techo amplio. No representa el consumo actual de un hogar bogotano ni sustituye sus facturas.")
    sf.add_table(
        doc,
        ["Magnitud", "Valor", "Cálculo / origen"],
        [
            ("Demanda anual D", "1.872 kWh", "12 × 156 kWh"),
            ("Rendimiento anual Y", "1.260,65 kWh/kWp", "PVGIS 5.3"),
            ("Potencia por panel P", "0,450 kWp", "Ficha LONGi"),
            ("Necesidad continua n*", "3,300", "1.872 /(0,450 × 1.260,65)"),
            ("Recomendación N", "4 paneles", "ceil(3,300)"),
            ("Potencia nominal", "1,80 kWp", "4 × 0,450"),
            ("Producción anual", "2.269,17 kWh", "1,80 × 1.260,65"),
            ("Cobertura anual", "121,22 %", "2.269,17 /1.872"),
            ("Máximo de techo", "9 paneles", "⌊20 /(1,722×1,134×1,10)⌋"),
        ],
        [2600, 2100, 4660],
        font_size=9.2,
    )
    sf.add_callout(doc, "Decisión", "Cuatro paneles compensan la meta anual en el modelo y caben por área. Sin embargo, junio llega a 98,17 % mensual; para alcanzar 100 % en todos los meses modelados se requieren 5 paneles. El usuario debe elegir el criterio.")

    sf.add_section_title(doc, "8. Contraste entre simuladores e incertidumbre", page_break=True)
    sf.add_table(
        doc,
        ["Ciudad", "PVGIS", "PVWatts", "Diferencia anual", "Decisión entera"],
        [
            (c["city"], f"{c['pvgis']:.2f}", f"{c['pvwatts']:.2f}", f"{c['annual_diff']:+.2f} %", f"{c['panels']} / {c['pvwatts_panels']}")
            for c in CASES
        ],
        [1600, 1500, 1500, 2150, 2610],
        font_size=9.0,
    )
    add_figure(doc, ASSET_DIR / "yield-comparison.png", "Figura 2. Contraste intermodelo del rendimiento específico anual.", "Comparación anual de PVGIS y PVWatts para cinco ubicaciones colombianas.")
    sf.add_para(doc, "Las diferencias absolutas anuales abarcan 0,82 %–20,64 %. Cali y Quibdó superan 10 %, pero solo Quibdó cambia la decisión entera: 4 paneles con PVGIS y 3 con PVWatts. Esto demuestra por qué un reporte responsable debe conservar proveedor, versión, base meteorológica y rango, en vez de ofrecer un único número como certeza.")
    sf.add_callout(doc, "Sin validación experimental", "No se comparó contra generación medida porque no se dispone de un sistema real comparable. Las cinco ejecuciones permiten observar el comportamiento matemático del modelo, comprobar el software en casos programados y hacer visibles las discrepancias entre simuladores: son especialmente altas en Cali y Quibdó, y solo Quibdó cambia la decisión entera. No demuestran exactitud en campo, ahorro real ni desempeño del techo.", fill=sf.PALE_GOLD, accent=sf.GOLD)

    sf.add_section_title(doc, "9. Por qué las simulaciones son esenciales")
    sf.add_para(doc, "Una instalación física exige capital, permisos, condiciones seguras y al menos varios meses de medición. Para un curso semestral, la simulación permite explorar condiciones contrastantes, repetir los cálculos y auditar las transformaciones. Esta evidencia sirve para la verificación interna del modelo y del software; la validación experimental queda pendiente hasta disponer de una instalación y producción medida.")
    sf.add_table(
        doc,
        ["Sí demuestran", "No demuestran"],
        [
            ("Que las ecuaciones producen resultados consistentes", "Que una estructura soporte el sistema"),
            ("Que n′(Y) cuantifica cómo el recurso cambia la necesidad continua", "Que una factura llegue a cero"),
            ("Que el techo inviable se reporta sin ocultar la necesidad", "Que no haya sombras u obstáculos"),
            ("Que dos proveedores pueden cambiar una decisión", "Que un proveedor sea la verdad física"),
            ("Que el proceso es repetible con entradas archivadas", "Que exista disposición a pagar o product–market fit"),
        ],
        [4680, 4680],
        font_size=9.1,
    )

    sf.add_section_title(doc, "10. Conclusiones")
    conclusions = [
        "SolarFit responde la pregunta académica con una potencia continua y una cantidad entera, sin mezclar kW con kWh.",
        "La derivada n′(Y) hace explicable la recomendación: su signo negativo y magnitud relacionan directamente recurso solar y necesidad continua.",
        "El balance anual y el criterio conservador resuelven preguntas distintas; el producto no debe ocultar esa decisión.",
        "Cuatro casos satisfacen la meta anual con el área supuesta; Quibdó activa correctamente una advertencia de inviabilidad por área.",
        "PVGIS y PVWatts coinciden en orden de magnitud, pero las diferencias de Cali y Quibdó impiden presentar un único resultado como certeza.",
        "El MVP es una herramienta de prefactibilidad explicable; la fase siguiente debe validar facturas, techos, propuestas profesionales y producción medida.",
    ]
    for item in conclusions:
        sf.add_bullet(doc, item, 0)

    sf.add_section_title(doc, "11. Escalamiento: de MVP a startup", page_break=True)
    sf.add_para(doc, "La siguiente versión no debe limitarse a nuevas pantallas. Debe cerrar el flujo factura → techo → producción → economía → instalador → conexión → monitoreo. La tesis recomendada es B2B2C: prefactibilidad gratuita para hogares y herramientas pagas para instaladores, financiadores o comercializadores.")
    sf.add_table(
        doc,
        ["Fase", "Producto", "Puerta de avance"],
        [
            ("0. Confianza técnica (0–2 meses)", "Carga de facturas, doble modelo, reporte trazable y privacidad", "Cálculos con fuente/versión y cero errores críticos"),
            ("1. Piloto concierge (2–5 meses)", "50–60 hogares + 10–15 instaladores", "Discrepancia mediana ≤1 panel o ≤10 %; ≥25 % pide cotización"),
            ("2. SolarFit v1 B2B2C (5–12 meses)", "OCR, techo, tarifa, economía, CRM y propuestas comparables", "Conversión lead–visita ≥20 % y cotización–venta ≥10 %"),
            ("3. Transacción y monitoreo (12–24 meses)", "Financiación aliada, seguimiento OR e inversor", "≥100 sistemas; error de ahorro medido ≤15 %"),
            ("4. Plataforma (24–36 meses)", "API, white-label, baterías y comunidades", "LTV/CAC ≥3 y dos canales repetibles"),
        ],
        [2300, 3860, 3200],
        font_size=8.6,
    )
    sf.add_para(doc, "Los umbrales son objetivos internos recomendados, no exigencias regulatorias. Antes de vender, SolarFit debe versionar reglas CREG, mantener el informe como prefactibilidad, aplicar RETIE y proteger los datos de facturas. Véanse [7][8].")

    sf.add_section_title(doc, "12. Limitaciones y riesgos")
    for text in (
        "La demanda usa promedios oficiales agregados repetidos cada mes; no refleja una vivienda, estacionalidad ni simultaneidad horaria.",
        "El área útil es un escalar y no sustituye un plano, visita o revisión estructural.",
        "PVGIS y PVWatts no capturan todas las sombras, suciedad, indisponibilidad o decisiones de montaje.",
        "La cobertura energética no equivale a ahorro monetario ni factura cero.",
        "La comparación de proveedores no es calibración contra una instalación medida.",
        "Las APIs y la regulación pueden cambiar; fuente, fecha y versión deben conservarse.",
    ):
        sf.add_bullet(doc, text, 0)

    add_identity_and_ai(doc)
    sf.add_section_title(doc, "Plan del pitch de 3 minutos", page_break=True)
    sf.add_table(
        doc,
        ["Tiempo", "Pregunta", "Mensaje"],
        [
            ("0:00–0:35", "¿Qué problema?", "Consumo y potencia no bastan; importan mes, clima, pérdidas y techo."),
            ("0:35–1:25", "¿Cómo se resolvió?", "Balance energético, n(Y), n′(Y), techo separado, MVP y APIs oficiales."),
            ("1:25–2:25", "¿Qué respalda el resultado y cuáles son sus límites?", "Cinco simulaciones, 38 pruebas y contraste PVGIS–PVWatts; no hay validación experimental."),
            ("2:25–3:00", "¿Qué aprendimos?", "Quibdó activa una advertencia de inviabilidad por restricción de techo; el siguiente paso es un piloto medido."),
        ],
        [1500, 2300, 5560],
        font_size=9.2,
    )
    sf.add_callout(doc, "Regla de la rúbrica", "Durante el pitch el único apoyo es el póster; no se debe presentar una secuencia de diapositivas.")

    sf.add_section_title(doc, "Referencias", page_break=True)
    add_references(doc)
    sf.add_section_title(doc, "Anexo A · Series de consumo simuladas", page_break=True)
    sf.add_table(
        doc,
        ["Mes"] + [c["city"] for c in CASES],
        [(month, *[str(c["consumption"][i]) for c in CASES]) for i, month in enumerate(("Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"))],
        [1050, 1660, 1660, 1660, 1660, 1670],
        font_size=8.8,
    )
    sf.add_table_source(doc, "Unidad: kWh/mes. Valores calibrados con promedios agregados SUI 2018 [5] y repetidos durante los doce meses porque la fuente no publica estacionalidad por caso. No son facturas ni mediciones de viviendas.")
    sf.add_section_title(doc, "Anexo B · Checklist de entrega", page_break=True)
    for item in (
        "[ ] Completar nombres, códigos y contribuciones reales en Word, póster y registro de IA.",
        "[ ] Confirmar la fecha y el límite de páginas en la tarea activa de AVATA.",
        "[ ] Mantener la expresión ‘casos simulados calibrados’; no escribir ‘facturas reales’.",
        "[ ] Adjuntar o conservar el libro de validación y la carpeta de evidencia.",
        "[ ] Verificar que póster, pitch e informe usan las mismas cifras.",
        "[ ] Ensayar el pitch a 3 minutos usando únicamente el póster.",
        "[ ] Asegurar que cada integrante pueda explicar el balance, n(Y), n′(Y), el algoritmo, los resultados y las limitaciones.",
    ):
        sf.add_para(doc, item, align=WD_ALIGN_PARAGRAPH.LEFT, after=5)
    finalize_doc(doc)
    return doc


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    build_yield_chart(ASSET_DIR / "yield-comparison.png")
    build_coverage_chart(ASSET_DIR / "coverage-cases.png")
    hito2 = build_hito2()
    hito2.save(HITO2_OUTPUT)
    final = build_final()
    final.save(FINAL_OUTPUT)
    print(f"Created {HITO2_OUTPUT}")
    print(f"Created {FINAL_OUTPUT}")


if __name__ == "__main__":
    main()
